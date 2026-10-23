#!/usr/bin/env python3
"""Exporta o relatório de alocações KWMS usando o login por e-mail e a VPN."""

from __future__ import annotations

import argparse
import ctypes
import logging
import math
import os
import shutil
import sys
import tempfile
import time
import unicodedata
from logging.handlers import RotatingFileHandler
from pathlib import Path
from typing import Any, Callable, Sequence
from urllib.parse import urlparse

from pipefy_clientes_tool_edge import FileLock, build_edge


LOGIN_URL = "https://kwms.br.kworld.kpmg.com/wms/login"
REPORTING_URL = "https://kwms.br.kworld.kpmg.com/wms/reporting"
FINAL_FILENAME = "Detalhes de Alocação_- Planos Aprovados e Baixados.xlsx"
DEFAULT_DESTINATION = r"C:\Users\chailer\OneDrive - KPMG\BR - CAIO 2-1 Gestão - Dashboard\Planilhas de Apoio\KWMS"
DEFAULT_EMAIL = "chailer@kpmg.com.br"
VPN_SKIP_EXIT_CODE = 20
LOGGER = logging.getLogger("kwms_alocacoes")


class VpnRequiredError(RuntimeError):
    pass


def is_empty_response(value: str) -> bool:
    text = value.casefold().replace("’", "'")
    return "err_empty_response" in text or (
        "kwms.br.kworld.kpmg.com" in text and "didn't send any data" in text
    )


def check_vpn_error(driver: Any, error: Exception | None = None) -> None:
    if error is not None and is_empty_response(str(error)):
        raise VpnRequiredError("ERR_EMPTY_RESPONSE: ative a VPN da KPMG para acessar o KWMS.") from error
    try:
        page = driver.execute_script("return document.body ? document.body.innerText : '';") or ""
    except Exception:
        page = ""
    if is_empty_response(str(page)):
        raise VpnRequiredError("ERR_EMPTY_RESPONSE: ative a VPN da KPMG para acessar o KWMS.")


def navigate(driver: Any, url: str) -> None:
    try:
        driver.get(url)
    except Exception as exc:
        check_vpn_error(driver, exc)
        raise
    check_vpn_error(driver)


def wait_until(driver: Any, predicate: Callable[[Any], Any], timeout: int, description: str) -> Any:
    from selenium.common.exceptions import StaleElementReferenceException

    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        check_vpn_error(driver)
        try:
            result = predicate(driver)
            if result:
                return result
        except StaleElementReferenceException:
            pass
        time.sleep(0.5)
    check_vpn_error(driver)
    raise TimeoutError(f"O KWMS não apresentou {description} em {timeout}s. Verifique o login e a disponibilidade do sistema.")


def visible_elements(driver: Any, selector: str) -> list[Any]:
    return [element for element in driver.find_elements("css selector", selector)
            if element.is_displayed() and element.is_enabled()]


def login_state(driver: Any) -> Any:
    emails = visible_elements(driver, "input[type='email']")
    if emails:
        return "email", emails[0]
    url = urlparse(driver.current_url)
    if url.hostname == "kwms.br.kworld.kpmg.com" and url.path.startswith("/wms") and url.path.rstrip("/") != "/wms/login":
        return "authenticated", None
    return False


def open_reporting(driver: Any, email: str, timeout: int) -> Any:
    navigate(driver, LOGIN_URL)
    state, element = wait_until(driver, login_state, timeout, "o campo de e-mail ou a sessão autenticada")
    if state == "email":
        element.clear()
        element.send_keys(email)
        buttons = wait_until(driver, lambda current: visible_elements(current, "button[type='submit'][aria-label='Login with email']"), timeout, "o botão Continue")
        buttons[0].click()
        def authenticated(current):
            result = login_state(current)
            return result if result and result[0] == "authenticated" else False
        wait_until(driver, authenticated,
                   timeout, "a conclusão do login por e-mail")
    navigate(driver, REPORTING_URL)
    buttons = wait_until(driver, lambda current: visible_elements(current, "button[title='Baixar arquivo']"), timeout, "o botão Baixar arquivo do relatório")
    if len(buttons) != 1:
        raise RuntimeError("O reporting mostrou vários botões 'Baixar arquivo'. É necessário identificar o relatório correto antes de baixar.")
    return buttons[0]


def flash_window(hwnd: int) -> None:
    if os.name != "nt" or not hwnd:
        return
    from ctypes import wintypes

    class FlashInfo(ctypes.Structure):
        _fields_ = [("cbSize", wintypes.UINT), ("hwnd", wintypes.HWND),
                    ("dwFlags", wintypes.DWORD), ("uCount", wintypes.UINT),
                    ("dwTimeout", wintypes.DWORD)]

    user32 = ctypes.windll.user32
    user32.FlashWindowEx.argtypes = [ctypes.POINTER(FlashInfo)]
    user32.FlashWindowEx.restype = wintypes.BOOL
    info = FlashInfo(ctypes.sizeof(FlashInfo), hwnd, 3, 6, 500)
    user32.FlashWindowEx(ctypes.byref(info))


def flash_console() -> None:
    if os.name == "nt":
        from ctypes import wintypes
        kernel32 = ctypes.windll.kernel32
        kernel32.GetConsoleWindow.restype = wintypes.HWND
        flash_window(kernel32.GetConsoleWindow())


def notify_vpn_and_wait(seconds: int) -> None:
    LOGGER.warning("Ative a VPN da KPMG. A segunda e última tentativa do KWMS ocorrerá em %s segundos.", seconds)
    deadline = time.monotonic() + seconds
    root = None
    try:
        flash_console()
        import tkinter as tk
        from tkinter import ttk
        root = tk.Tk()
        root.title("KWMS — ative a VPN")
        root.geometry("560x220")
        root.resizable(False, False)
        frame = ttk.Frame(root, padding=22)
        frame.pack(fill="both", expand=True)
        ttk.Label(frame, text="Ative a VPN da KPMG", font=("Segoe UI", 16, "bold")).pack(anchor="w")
        ttk.Label(frame, text="O KWMS retornou ERR_EMPTY_RESPONSE.\nLigue a VPN para permitir a atualização do relatório.", font=("Segoe UI", 10)).pack(anchor="w", pady=12)
        countdown = ttk.Label(frame, font=("Segoe UI", 11, "bold"))
        countdown.pack(anchor="w")
        ttk.Label(frame, text="Se falhar novamente, o lote seguirá para a próxima automação.").pack(anchor="w", pady=10)
        root.protocol("WM_DELETE_WINDOW", root.iconify)
        root.update_idletasks()
        if os.name == "nt":
            from ctypes import wintypes
            user32 = ctypes.windll.user32
            user32.GetAncestor.argtypes = [wintypes.HWND, wintypes.UINT]
            user32.GetAncestor.restype = wintypes.HWND
            hwnd = user32.GetAncestor(root.winfo_id(), 2)
            flash_window(hwnd)
        def tick():
            remaining = math.ceil(deadline - time.monotonic())
            if remaining <= 0:
                root.destroy()
                return
            countdown.configure(text=f"Nova tentativa em {remaining} segundos.")
            root.after(min(500, max(1, int((deadline - time.monotonic()) * 1000))), tick)
        tick()
        root.mainloop()
    except Exception as exc:
        LOGGER.warning("Não foi possível exibir o aviso visual de VPN: %s", exc)
        if root is not None:
            try:
                root.destroy()
            except Exception:
                pass
    remaining = deadline - time.monotonic()
    if remaining > 0:
        time.sleep(remaining)


def with_vpn_retry(operation: Callable[[], Any], wait: Callable[[int], None] | None = None) -> Any:
    wait = wait or notify_vpn_and_wait
    for attempt in (1, 2):
        try:
            return operation()
        except VpnRequiredError:
            if attempt == 2:
                LOGGER.warning("KWMS ignorado: a segunda tentativa também retornou ERR_EMPTY_RESPONSE. Ative a VPN. Continuando para a próxima automação.")
                raise
            wait(30)
    raise AssertionError("Fluxo de tentativas inválido.")


def matches_download_name(name: str) -> bool:
    normalized = unicodedata.normalize("NFKC", name).casefold()
    prefix = unicodedata.normalize("NFKC", Path(FINAL_FILENAME).stem).casefold()
    return normalized.endswith(".xlsx") and normalized.startswith(prefix)


def wait_for_download(directory: Path, timeout: int, driver: Any) -> Path:
    deadline = time.monotonic() + timeout
    previous: dict[Path, tuple[int, int]] = {}
    while time.monotonic() < deadline:
        check_vpn_error(driver)
        files = [path for path in directory.iterdir() if path.is_file()]
        partials = any(path.name.lower().endswith((".crdownload", ".tmp", ".part")) for path in files)
        candidates = [path for path in files if matches_download_name(path.name)]
        if not partials and len(candidates) == 1:
            candidate = candidates[0]
            try:
                stat = candidate.stat()
            except FileNotFoundError:
                time.sleep(1)
                continue
            signature = stat.st_size, stat.st_mtime_ns
            if stat.st_size > 0 and previous.get(candidate) == signature:
                return candidate
            previous[candidate] = signature
        elif len(candidates) > 1:
            raise RuntimeError("Foram baixados vários relatórios KWMS; o arquivo anterior foi preservado.")
        elif not partials and any(path.suffix.lower() == ".xlsx" for path in files):
            raise RuntimeError("O XLSX baixado não possui o nome esperado do relatório de alocações; o arquivo anterior foi preservado.")
        time.sleep(1)
    raise TimeoutError(f"Download KWMS não concluído em {timeout}s. O arquivo anterior foi preservado.")


def validate_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            populated_rows = 0
            for sheet in workbook.worksheets:
                sheet.reset_dimensions()
                for row in sheet.iter_rows(values_only=True):
                    if any(value is not None and str(value).strip() for value in row):
                        populated_rows += 1
            if populated_rows < 2:
                raise ValueError("A exportação está vazia ou contém apenas uma linha de cabeçalho.")
        finally:
            workbook.close()
    except Exception as exc:
        raise ValueError(f"Relatório KWMS inválido ou vazio: {exc}. O arquivo anterior foi preservado.") from exc


def publish_workbook(source: Path, destination: Path) -> Path:
    validate_workbook(source)
    destination.mkdir(parents=True, exist_ok=True)
    target = destination / FINAL_FILENAME
    descriptor, name = tempfile.mkstemp(prefix=".kwms-", suffix=".xlsx.tmp", dir=destination)
    os.close(descriptor)
    temporary = Path(name)
    try:
        shutil.copyfile(source, temporary)
        with temporary.open("r+b") as handle:
            handle.flush()
            os.fsync(handle.fileno())
        try:
            os.replace(temporary, target)
        except PermissionError as exc:
            raise PermissionError(f"Feche o relatório KWMS no Excel para substituir {target}. O arquivo anterior foi preservado.") from exc
    finally:
        temporary.unlink(missing_ok=True)
    return target


def execute(args: argparse.Namespace) -> int:
    work = Path(args.work_dir).expanduser().resolve()
    work.mkdir(parents=True, exist_ok=True)
    download = Path(tempfile.mkdtemp(prefix="run-", dir=work))
    driver = None
    success = False
    try:
        with FileLock(work / "kwms.lock", timeout=30):
            driver = build_edge(headless=args.headless, user_data_dir=args.user_data_dir,
                                download_dir=download, browser_start_timeout=120)
            driver.set_page_load_timeout(args.page_timeout)
            def operation():
                try:
                    attempt_dir = Path(tempfile.mkdtemp(prefix="attempt-", dir=download))
                    params = {"behavior": "allow", "downloadPath": str(attempt_dir), "eventsEnabled": True}
                    try:
                        driver.execute_cdp_cmd("Browser.setDownloadBehavior", params)
                    except Exception:
                        driver.execute_cdp_cmd("Page.setDownloadBehavior", params)
                    button = open_reporting(driver, args.email, args.page_timeout)
                    if args.command == "login":
                        LOGGER.info("Login KWMS concluído com o e-mail informado.")
                        return
                    button.click()
                    check_vpn_error(driver)
                    source = wait_for_download(attempt_dir, args.download_timeout, driver)
                    target = publish_workbook(source, Path(args.destination).expanduser().resolve())
                    LOGGER.info("Relatório KWMS atualizado: %s", target)
                except VpnRequiredError:
                    raise
                except Exception as exc:
                    check_vpn_error(driver, exc)
                    raise
            with_vpn_retry(operation)
            success = True
            return 0
    except VpnRequiredError:
        return VPN_SKIP_EXIT_CODE
    finally:
        if driver is not None:
            try:
                driver.quit()
            except Exception as exc:
                LOGGER.warning("Não foi possível fechar a sessão KWMS: %s", exc)
        if success:
            shutil.rmtree(download, ignore_errors=True)
        else:
            LOGGER.warning("Pasta de diagnóstico do KWMS: %s", download)


def build_parser() -> argparse.ArgumentParser:
    extraction = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(description="Exporta o relatório KWMS e substitui o XLSX de apoio.")
    parser.add_argument("--log-file", default=str(extraction / "logs" / "kwms_alocacoes_tool_edge.log"))
    parser.add_argument("--email", default=os.environ.get("KWMS_EMAIL", DEFAULT_EMAIL))
    parser.add_argument("--user-data-dir", default=os.environ.get("KWMS_EDGE_PROFILE", r"C:\SeleniumEdgeProfile"))
    parser.add_argument("--destination", default=os.environ.get("KWMS_OUTPUT_DIR", DEFAULT_DESTINATION))
    parser.add_argument("--work-dir", default=str(Path(os.environ.get("PIPEFY_WORK_DIR", str(extraction / ".pipefy_orchestrator" / "downloads"))) / "kwms"))
    parser.add_argument("--page-timeout", type=int, default=60)
    parser.add_argument("--download-timeout", type=int, default=600)
    commands = parser.add_subparsers(dest="command", required=True)
    for command in ("run", "login"):
        subparser = commands.add_parser(command)
        subparser.add_argument("--headless", action="store_true")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    log_file = Path(args.log_file).expanduser().resolve()
    log_file.parent.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s",
                        datefmt="%H:%M:%S", handlers=[logging.StreamHandler(),
                        RotatingFileHandler(log_file, maxBytes=5_000_000, backupCount=3, encoding="utf-8")])
    try:
        if args.page_timeout <= 0 or args.download_timeout <= 0:
            raise ValueError("Os tempos limite devem ser positivos.")
        return execute(args)
    except KeyboardInterrupt:
        LOGGER.warning("Rotina KWMS interrompida pelo usuário.")
        return 130
    except Exception as exc:
        LOGGER.error("Falha no KWMS: %s", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
