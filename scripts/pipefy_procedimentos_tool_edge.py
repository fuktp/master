#!/usr/bin/env python3
"""Exporta Procedimentos e publica procedimentos/Procedimentos.xlsx no PBi."""

from __future__ import annotations

import argparse
import base64
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from dataclasses import dataclass
from datetime import datetime
from logging.handlers import RotatingFileHandler
from pathlib import Path
from typing import Any, Sequence
from urllib.parse import urlparse


APP_NAME = "pipefy_procedimentos_tool"
BASE_DIR = Path(__file__).resolve().parent
EXTRACTION_DIR = BASE_DIR.parent
DEFAULT_OUTPUT_DIR = Path(
    os.environ.get(
        "PIPEFY_OUTPUT_DIR",
        r"C:\Users\chailer\OneDrive - KPMG\BR - CAIO 2-1 Gestão - Dashboard\Testes - PBi",
    )
).expanduser()
DEFAULT_WORK_DIR = Path(
    os.environ.get(
        "PIPEFY_WORK_DIR",
        str(EXTRACTION_DIR / ".pipefy_orchestrator" / "downloads"),
    )
).expanduser()
EDGE_DRIVER_PATH = Path(
    os.environ.get(
        "PIPEFY_EDGE_DRIVER_PATH",
        r"C:\Users\chailer\Downloads\edgedriver_win64\msedgedriver.exe",
    )
)
# Edge e msedgedriver podem ser atualizados para qualquer versão principal
# igual ou superior sem exigir alteração nesta constante.
MIN_EDGE_MAJOR_VERSION = 150
DEFAULT_EDGE_PROFILE = r"C:\SeleniumEdgeProfile"
DEFAULT_EMAIL = "chailer@kpmg.com.br"
DEFAULT_DATABASE_URL = "https://app.pipefy.com/apollo_databases/303446333"
DEFAULT_PROCEDURES_DIR = str(DEFAULT_OUTPUT_DIR / "procedimentos")
DEFAULT_PROCEDURES_STAGING_DIR = str(DEFAULT_WORK_DIR / "procedimentos")
FINAL_FILENAME = "Procedimentos.xlsx"
DEFAULT_LOG_FILE = str(EXTRACTION_DIR / "logs" / "pipefy_procedimentos_tool.log")
DEFAULT_TASK_NAME = "Pipefy Procedimentos Export"
# Procedimentos roda dez minutos depois de Eventos em cada rodada útil.
TASK_TIMES = ("09:10", "11:10", "13:10", "15:10", "17:10")
TASK_WEEKDAYS = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday")

EXPECTED_COLUMNS = (
    "Código",
    "Nome",
    "Ticket KDSC / Audit Tech",
    "Modelo de WP / Template MESP",
    "Soluções digitais aplicáveis",
    "Sugestão de Prompt:",
    "Descrição do Procedimento:",
    "RMMs aplicáveis",
    "Etapas do procedimento:",
    "Documentos sugeridos:",
    "Processo",
    "Área Final",
    "Cargo sugerido para execução do procedimento",
    "Tempo estimado:",
    "Catálogo",
    "ID",
    "Title",
    "Status",
    "Creator",
    "Finished at",
    "Created at",
    "Updated at",
    "Expired at",
    "Overdue",
)
# Renomeação conhecida no Pipefy. Os aliases valem somente na validação;
# o XLSX publicado conserva os cabeçalhos e os dados da exportação original.
COLUMN_ALIASES = {
    "Ticket KDSC": "Ticket KDSC / Audit Tech",
}
EXPORT_LOCATORS = (
    ("css selector", "button[aria-label='Export'], button[aria-label='Exportar']"),
    (
        "xpath",
        "//button[.//*[normalize-space()='Export' or normalize-space()='Exportar'] "
        "or normalize-space()='Export' or normalize-space()='Exportar']",
    ),
)
SSO_SELECTOR = (
    "a[id*='social-KPMGOficial'][id*='saml'], "
    "a[href*='KPMGOficial'][href*='saml']"
)
SENSITIVE_HEADERS = {
    "authorization",
    "cookie",
    "set-cookie",
    "x-csrf-token",
    "x-xsrf-token",
    "x-pipefy-token",
    "x-pipefy-app-token",
    "access-token",
    "client",
    "uid",
}
LOGGER = logging.getLogger(APP_NAME)


def resolve_path(value: str | Path) -> Path:
    path = Path(value).expanduser()
    return path.resolve() if path.is_absolute() else (BASE_DIR / path).resolve()


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def configure_logging(verbose: bool, log_file: str, silent: bool = False) -> None:
    LOGGER.handlers.clear()
    LOGGER.propagate = False
    LOGGER.setLevel(logging.DEBUG if verbose else logging.INFO)
    formatter = logging.Formatter(
        "%(asctime)s | %(levelname)s | %(message)s",
        datefmt="%H:%M:%S",
    )

    if not silent:
        console = logging.StreamHandler()
        console.setFormatter(formatter)
        LOGGER.addHandler(console)

    path = resolve_path(log_file)
    path.parent.mkdir(parents=True, exist_ok=True)
    file_handler = RotatingFileHandler(
        path,
        maxBytes=5_000_000,
        backupCount=5,
        encoding="utf-8",
    )
    file_handler.setFormatter(formatter)
    LOGGER.addHandler(file_handler)


class FileLock:
    """Lock simples entre processos, compatível com Windows e Unix."""

    def __init__(self, path: Path, timeout: float = 30.0) -> None:
        self.path = path
        self.timeout = timeout
        self.handle: Any = None

    def __enter__(self) -> "FileLock":
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.handle = open(self.path, "a+b")
        self.handle.seek(0, os.SEEK_END)
        if self.handle.tell() == 0:
            self.handle.write(b"0")
            self.handle.flush()

        deadline = time.monotonic() + self.timeout
        while True:
            try:
                self.handle.seek(0)
                if os.name == "nt":
                    import msvcrt

                    msvcrt.locking(self.handle.fileno(), msvcrt.LK_NBLCK, 1)
                else:
                    import fcntl

                    fcntl.flock(self.handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
                return self
            except (OSError, BlockingIOError):
                if time.monotonic() >= deadline:
                    self.handle.close()
                    self.handle = None
                    raise TimeoutError(
                        f"Outra execução está usando o recurso: {self.path}"
                    )
                time.sleep(0.25)

    def __exit__(self, exc_type: Any, exc: Any, traceback: Any) -> None:
        if not self.handle:
            return
        try:
            self.handle.seek(0)
            if os.name == "nt":
                import msvcrt

                msvcrt.locking(self.handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl

                fcntl.flock(self.handle.fileno(), fcntl.LOCK_UN)
        finally:
            self.handle.close()
            self.handle = None


class ProfileInUseError(RuntimeError):
    pass


def _is_valid_profile_json(data: bytes) -> bool:
    try:
        value = json.loads(
            data.decode("utf-8-sig"),
            parse_constant=lambda token: (_ for _ in ()).throw(
                ValueError(f"constante JSON inválida: {token}")
            ),
        )
    except (UnicodeError, ValueError):
        return False
    return isinstance(value, dict)


def _invalid_profile_backup(path: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    return path.with_name(f"{path.name}.invalid-{timestamp}.bak")


def repair_invalid_edge_profile_json(profile: Path) -> tuple[Path, ...]:
    """Isola Preferences/Local State inválidos sem apagar o perfil ou cookies."""
    repaired: list[Path] = []
    candidates = (
        profile / "Default" / "Preferences",
        profile / "Local State",
    )
    for path in candidates:
        if not path.is_file():
            continue
        try:
            first = path.read_bytes()
        except OSError as exc:
            raise ProfileInUseError(
                f"Não foi possível ler o arquivo de perfil {path}: {exc}"
            ) from exc
        if _is_valid_profile_json(first):
            continue

        time.sleep(0.25)
        try:
            second = path.read_bytes()
        except FileNotFoundError:
            continue
        except OSError as exc:
            raise ProfileInUseError(
                f"Não foi possível reler o arquivo de perfil {path}: {exc}"
            ) from exc
        if _is_valid_profile_json(second):
            continue
        if first != second:
            raise ProfileInUseError(
                f"O perfil está sendo alterado por outro processo: {path}"
            )

        backup = _invalid_profile_backup(path)
        try:
            os.replace(path, backup)
        except OSError as exc:
            raise ProfileInUseError(
                f"Não foi possível isolar o JSON inválido {path}: {exc}"
            ) from exc
        repaired.append(backup)
        LOGGER.warning(
            "JSON inválido do perfil isolado: %s. Backup: %s",
            path,
            backup,
        )
    return tuple(repaired)


class PipefySessionExpiredError(RuntimeError):
    """Indica que uma execução headless encontrou a tela de login."""


def _command_output(command: Sequence[str]) -> str:
    creation_flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
    result = subprocess.run(
        list(command),
        capture_output=True,
        text=True,
        timeout=20,
        check=False,
        creationflags=creation_flags,
    )
    return "\n".join(part for part in (result.stdout, result.stderr) if part).strip()


def _major_version(version: str, component: str) -> int:
    """Extrai a versão principal de um componente Edge."""

    match = re.match(r"\s*(\d+)", version)
    if not match:
        raise RuntimeError(
            f"Não foi possível identificar a versão principal de {component}: "
            f"{version or '<vazia>'}."
        )
    return int(match.group(1))


def _validate_minimum_edge_version(version: str, component: str) -> int:
    """Aceita qualquer versão do componente a partir da mínima configurada."""

    major = _major_version(version, component)
    if major < MIN_EDGE_MAJOR_VERSION:
        raise RuntimeError(
            f"{component} incompatível: é necessária a versão "
            f"{MIN_EDGE_MAJOR_VERSION} ou superior; encontrada {version}."
        )
    return major


def validate_edge_driver() -> str:
    if not EDGE_DRIVER_PATH.is_file():
        raise FileNotFoundError(
            f"msedgedriver v{MIN_EDGE_MAJOR_VERSION} ou superior não encontrado "
            "no caminho configurado: "
            f"{EDGE_DRIVER_PATH}"
        )
    output = _command_output([str(EDGE_DRIVER_PATH), "--version"])
    match = re.search(r"\b(\d+\.\d+\.\d+\.\d+)\b", output)
    if not match:
        raise RuntimeError(
            f"Não foi possível identificar a versão de {EDGE_DRIVER_PATH}. "
            f"Saída: {output or '<vazia>'}"
        )
    version = match.group(1)
    _validate_minimum_edge_version(version, "msedgedriver")
    return version


def _build_edge_once(
    *,
    headless: bool,
    user_data_dir: str,
    download_dir: Path | None,
    capture_network: bool,
    start_minimized: bool,
) -> Any:
    from selenium import webdriver
    from selenium.webdriver.edge.options import Options
    from selenium.webdriver.edge.service import Service

    driver_version = validate_edge_driver()
    profile = Path(user_data_dir).expanduser().resolve()
    profile.mkdir(parents=True, exist_ok=True)
    repaired_profile_files = repair_invalid_edge_profile_json(profile)
    if repaired_profile_files:
        LOGGER.info(
            "Perfil do Edge reparado; %s arquivo(s) inválido(s) isolado(s).",
            len(repaired_profile_files),
        )

    options = Options()
    options.add_argument(f"--user-data-dir={profile}")
    options.add_argument("--no-first-run")
    options.add_argument("--no-default-browser-check")
    options.add_argument("--disable-notifications")
    options.add_argument("--disable-background-mode")
    options.add_argument("--log-level=3")
    if headless:
        options.add_argument("--headless=new")
        options.add_argument("--window-size=1920,1080")
    elif start_minimized:
        options.add_argument("--start-minimized")
    else:
        options.add_argument("--start-maximized")

    prefs: dict[str, Any] = {
        "credentials_enable_service": False,
        "profile.password_manager_enabled": False,
    }
    if download_dir:
        download_dir.mkdir(parents=True, exist_ok=True)
        prefs.update(
            {
                "download.default_directory": str(download_dir),
                "download.prompt_for_download": False,
                "download.directory_upgrade": True,
                "safebrowsing.enabled": True,
            }
        )
    options.add_experimental_option("prefs", prefs)
    if capture_network:
        options.set_capability("ms:loggingPrefs", {"performance": "ALL"})

    LOGGER.info(
        "Iniciando Edge com msedgedriver %s (versão mínima aceita: v%s).",
        driver_version,
        MIN_EDGE_MAJOR_VERSION,
    )
    try:
        driver = webdriver.Edge(
            service=Service(str(EDGE_DRIVER_PATH), log_output=os.devnull),
            options=options,
        )
    except Exception as exc:
        detail = str(exc)
        if (
            "DevToolsActivePort" in detail
            or "user data directory is already in use" in detail.lower()
        ):
            raise ProfileInUseError(
                f"O perfil dedicado está em uso: {profile}"
            ) from exc
        raise RuntimeError(f"Falha ao iniciar Microsoft Edge: {detail}") from exc

    try:
        browser_version = str(driver.capabilities.get("browserVersion", ""))
        browser_major = _validate_minimum_edge_version(
            browser_version,
            "Microsoft Edge",
        )
        driver_major = _major_version(driver_version, "msedgedriver")

        if browser_major != driver_major:
            LOGGER.warning(
                "Versões principais diferentes; a execução continuará: "
                "Edge=%s, msedgedriver=%s. Se o próprio Selenium recusar a "
                "sessão, atualize o executável msedgedriver.",
                browser_version,
                driver_version,
            )
        else:
            LOGGER.info(
                "Versões aceitas: Edge=%s, msedgedriver=%s.",
                browser_version,
                driver_version,
            )

        if start_minimized:
            try:
                driver.minimize_window()
            except Exception as exc:
                LOGGER.warning("Não foi possível minimizar o Edge: %s", exc)

        if download_dir:
            params = {
                "behavior": "allow",
                "downloadPath": str(download_dir),
                "eventsEnabled": True,
            }
            try:
                driver.execute_cdp_cmd("Browser.setDownloadBehavior", params)
            except Exception:
                driver.execute_cdp_cmd("Page.setDownloadBehavior", params)

        if capture_network:
            driver.execute_cdp_cmd(
                "Network.enable",
                {
                    "maxTotalBufferSize": 100_000_000,
                    "maxResourceBufferSize": 10_000_000,
                    "maxPostDataSize": 10_000_000,
                },
            )
            if "performance" not in driver.log_types:
                raise RuntimeError(
                    "Edge iniciou, mas o log CDP 'performance' não está disponível."
                )
        return driver
    except Exception:
        driver.quit()
        raise


def build_edge(
    *,
    headless: bool,
    user_data_dir: str,
    download_dir: Path | None = None,
    capture_network: bool = False,
    browser_start_timeout: int = 900,
    start_minimized: bool = False,
) -> Any:
    deadline = time.monotonic() + browser_start_timeout
    while True:
        try:
            return _build_edge_once(
                headless=headless,
                user_data_dir=user_data_dir,
                download_dir=download_dir,
                capture_network=capture_network,
                start_minimized=start_minimized,
            )
        except ProfileInUseError:
            remaining = int(deadline - time.monotonic())
            if remaining <= 0:
                raise TimeoutError(
                    "O perfil do Edge permaneceu ocupado. Verifique se outro export "
                    "Pipefy ou uma janela manual ainda está em execução."
                )
            LOGGER.warning(
                "Perfil do Edge ocupado; nova tentativa em 15s (%ss restantes).",
                remaining,
            )
            time.sleep(min(15, remaining))


def _first_element(driver: Any, locators: Sequence[tuple[str, str]]) -> Any | None:
    from selenium.common.exceptions import StaleElementReferenceException

    for by, selector in locators:
        for element in driver.find_elements(by, selector):
            try:
                if element.is_displayed() and element.is_enabled():
                    return element
            except StaleElementReferenceException:
                continue
    return None


def _export_button(driver: Any) -> Any | None:
    return _first_element(driver, EXPORT_LOCATORS)


def _login_state(driver: Any) -> tuple[str, Any] | bool:
    button = _export_button(driver)
    if button:
        return "ready", button
    for element in driver.find_elements("id", "username"):
        if element.is_displayed():
            return "login", element
    return False


def _click(driver: Any, element: Any) -> None:
    try:
        driver.execute_script(
            "arguments[0].scrollIntoView({block: 'center'});", element
        )
        element.click()
    except Exception:
        driver.execute_script("arguments[0].click();", element)


def open_database(
    driver: Any,
    *,
    email: str,
    database_url: str,
    headless: bool,
    step_timeout: int,
    sso_timeout: int,
) -> Any:
    from selenium.common.exceptions import TimeoutException
    from selenium.webdriver.support.ui import WebDriverWait

    LOGGER.info("Abrindo tabela Procedimentos: %s", database_url)
    driver.get(database_url)
    wait = WebDriverWait(driver, step_timeout)
    try:
        state, element = wait.until(_login_state)
    except TimeoutException as exc:
        raise RuntimeError(
            "A página não exibiu nem o login nem o botão Export. "
            f"URL atual: {driver.current_url}"
        ) from exc

    if state == "ready":
        LOGGER.info("Sessão Pipefy já autenticada.")
        return element
    if headless:
        raise PipefySessionExpiredError(
            "A sessão SSO do Pipefy expirou durante a execução headless."
        )

    LOGGER.info("Preenchendo e-mail para autenticação SSO.")
    element.clear()
    element.send_keys(email)
    continue_button = wait.until(
        lambda current: _first_element(current, (("id", "kc-login"),))
    )
    _click(driver, continue_button)

    def sso_or_redirect(current: Any) -> tuple[str, Any | None] | bool:
        link = _first_element(current, (("css selector", SSO_SELECTOR),))
        if link:
            return "sso", link
        host = urlparse(current.current_url or "").hostname or ""
        if host == "app.pipefy.com" or (
            host and host != "pipefy.com" and not host.endswith(".pipefy.com")
        ):
            return "redirect", None
        return False

    try:
        next_state, sso_link = wait.until(sso_or_redirect)
        if next_state == "sso" and sso_link:
            LOGGER.info("Selecionando o provedor KPMGOficial.")
            _click(driver, sso_link)
    except TimeoutException as exc:
        raise RuntimeError("O provedor SSO KPMGOficial não foi localizado.") from exc

    LOGGER.info(
        "Aguardando SSO/MFA por até %ss; conclua na janela.", sso_timeout
    )
    try:
        WebDriverWait(driver, sso_timeout).until(
            lambda current: (
                urlparse(current.current_url or "").hostname == "app.pipefy.com"
            )
        )
    except TimeoutException as exc:
        raise RuntimeError("SSO/MFA não foi concluído dentro do tempo limite.") from exc

    driver.get(database_url)
    try:
        button = WebDriverWait(driver, step_timeout).until(_export_button)
    except TimeoutException as exc:
        raise RuntimeError(
            "Login concluído, mas o botão Export não apareceu na tabela Procedimentos."
        ) from exc
    LOGGER.info("Autenticação concluída e tabela Procedimentos pronta.")
    return button


def _perform_visible_login(
    *,
    email: str,
    database_url: str,
    user_data_dir: str,
    step_timeout: int,
    sso_timeout: int,
    browser_start_timeout: int,
    start_minimized: bool = False,
) -> None:
    """Executa o mesmo fluxo do comando ``login`` em uma janela do Edge."""

    driver = build_edge(
        headless=False,
        user_data_dir=user_data_dir,
        browser_start_timeout=browser_start_timeout,
        start_minimized=start_minimized,
    )
    try:
        open_database(
            driver,
            email=email,
            database_url=database_url,
            headless=False,
            step_timeout=step_timeout,
            sso_timeout=sso_timeout,
        )
    finally:
        driver.quit()


def _open_database_with_login_fallback(
    *,
    headless: bool,
    email: str,
    database_url: str,
    user_data_dir: str,
    step_timeout: int,
    sso_timeout: int,
    download_dir: Path | None = None,
    capture_network: bool = False,
    browser_start_timeout: int = 900,
) -> tuple[Any, Any]:
    """Abre o database e renova o SSO em janela visível quando necessário."""

    driver = build_edge(
        headless=headless,
        user_data_dir=user_data_dir,
        download_dir=download_dir,
        capture_network=capture_network,
        browser_start_timeout=browser_start_timeout,
    )
    try:
        button = open_database(
            driver,
            email=email,
            database_url=database_url,
            headless=headless,
            step_timeout=step_timeout,
            sso_timeout=sso_timeout,
        )
        return driver, button
    except PipefySessionExpiredError:
        driver.quit()
        if not headless:
            raise
    except Exception:
        driver.quit()
        raise

    LOGGER.warning(
        "Sessão Pipefy expirada. Fallback iniciado: abrindo o Edge visível "
        "para renovar o SSO/MFA."
    )
    _perform_visible_login(
        email=email,
        database_url=database_url,
        user_data_dir=user_data_dir,
        step_timeout=step_timeout,
        sso_timeout=sso_timeout,
        browser_start_timeout=browser_start_timeout,
        start_minimized=True,
    )

    LOGGER.info("Sessão Pipefy renovada. Repetindo a execução headless.")
    retry_driver = build_edge(
        headless=True,
        user_data_dir=user_data_dir,
        download_dir=download_dir,
        capture_network=capture_network,
        browser_start_timeout=browser_start_timeout,
    )
    try:
        button = open_database(
            retry_driver,
            email=email,
            database_url=database_url,
            headless=True,
            step_timeout=step_timeout,
            sso_timeout=sso_timeout,
        )
        return retry_driver, button
    except Exception:
        retry_driver.quit()
        raise


def _daily_login_paths(user_data_dir: str) -> tuple[Path, Path]:
    """Retorna o estado e o lock compartilhados pelos três scripts."""

    profile = Path(user_data_dir).expanduser().resolve()
    profile.mkdir(parents=True, exist_ok=True)
    return (
        profile / ".pipefy_daily_login.json",
        profile / ".pipefy_daily_login.lock",
    )


def _daily_login_completed(state_path: Path, today: str) -> bool:
    try:
        state = json.loads(state_path.read_text(encoding="utf-8"))
    except (OSError, TypeError, json.JSONDecodeError):
        return False
    return isinstance(state, dict) and state.get("date") == today


def _write_daily_login_state(state_path: Path, today: str) -> None:
    descriptor, name = tempfile.mkstemp(
        prefix=".pipefy_daily_login.",
        suffix=".json.tmp",
        dir=state_path.parent,
    )
    temporary = Path(name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as file:
            json.dump(
                {
                    "date": today,
                    "authenticated_at": now_iso(),
                    "source": APP_NAME,
                },
                file,
                ensure_ascii=False,
            )
            file.write("\n")
            file.flush()
            os.fsync(file.fileno())
        os.replace(temporary, state_path)
    finally:
        temporary.unlink(missing_ok=True)


def _mark_daily_login(user_data_dir: str) -> None:
    state_path, _ = _daily_login_paths(user_data_dir)
    today = datetime.now().astimezone().date().isoformat()
    _write_daily_login_state(state_path, today)
    LOGGER.info("Login diário do Pipefy confirmado em %s.", today)


def _ensure_daily_login(
    *,
    email: str,
    database_url: str,
    user_data_dir: str,
    step_timeout: int,
    sso_timeout: int,
    browser_start_timeout: int,
) -> None:
    """Executa uma vez ao dia o fluxo visível equivalente ao comando login."""

    state_path, lock_path = _daily_login_paths(user_data_dir)
    today = datetime.now().astimezone().date().isoformat()
    daily_lock_timeout = max(
        float(browser_start_timeout + sso_timeout + (step_timeout * 2) + 120),
        1800.0,
    )

    with FileLock(lock_path, timeout=daily_lock_timeout):
        if _daily_login_completed(state_path, today):
            LOGGER.info("Login diário do Pipefy já executado em %s.", today)
            return

        LOGGER.info(
            "Primeira execução do dia: executando o fluxo 'login' com janela "
            "visível antes de extrair as bases."
        )
        _perform_visible_login(
            email=email,
            database_url=database_url,
            user_data_dir=user_data_dir,
            step_timeout=step_timeout,
            sso_timeout=sso_timeout,
            browser_start_timeout=browser_start_timeout,
        )
        _write_daily_login_state(state_path, today)
        LOGGER.info("Login diário do Pipefy confirmado em %s.", today)


def _file_signature(path: Path) -> tuple[int, int]:
    stat = path.stat()
    return stat.st_mtime_ns, stat.st_size


def directory_snapshot(directory: Path) -> dict[str, tuple[int, int]]:
    return {
        path.name: _file_signature(path)
        for path in directory.iterdir()
        if path.is_file()
    }


def wait_for_download(
    directory: Path,
    before: dict[str, tuple[int, int]],
    timeout: int,
) -> Path:
    deadline = time.monotonic() + timeout
    stable: dict[str, tuple[int, int]] = {}
    while time.monotonic() < deadline:
        current: dict[Path, tuple[int, int]] = {}
        for path in directory.iterdir():
            try:
                if path.is_file():
                    current[path] = _file_signature(path)
            except FileNotFoundError:
                continue

        changed = [
            path
            for path, signature in current.items()
            if before.get(path.name) != signature
        ]
        partials = [
            path
            for path in changed
            if path.name.lower().endswith((".crdownload", ".tmp"))
        ]
        candidates = [path for path in changed if path.suffix.lower() == ".xlsx"]
        if candidates and not partials:
            candidate = max(candidates, key=lambda path: current[path][0])
            size = current[candidate][1]
            previous_size, checks = stable.get(candidate.name, (-1, 0))
            checks = checks + 1 if size == previous_size and size > 0 else 0
            stable[candidate.name] = (size, checks)
            if checks >= 1:
                return candidate
        time.sleep(1)

    partials = [
        path.name
        for path in directory.iterdir()
        if path.name.lower().endswith((".crdownload", ".tmp"))
    ]
    detail = f" Arquivos parciais: {partials}." if partials else ""
    raise TimeoutError(
        f"Download XLSX não terminou em {timeout}s no diretório {directory}.{detail}"
    )


def _sheet_value(value: str) -> str | int:
    return int(value) if re.fullmatch(r"\d+", value) else value


@dataclass(frozen=True)
class WorkbookValidation:
    sheet_name: str
    data_rows: int
    file_size: int


def validate_procedures_workbook(
    path: Path,
    *,
    sheet: str,
    header_row: int,
    allow_empty: bool,
) -> WorkbookValidation:
    from openpyxl import load_workbook

    if not path.is_file():
        raise FileNotFoundError(f"Arquivo XLSX não encontrado: {path}")
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"A exportação não é um arquivo .xlsx: {path}")

    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError(f"O arquivo baixado não é um XLSX válido: {path}") from exc

    try:
        selected = _sheet_value(sheet)
        if isinstance(selected, int):
            if selected < 0 or selected >= len(workbook.worksheets):
                raise ValueError(
                    f"Índice de aba inexistente: {selected}. "
                    f"Total de abas: {len(workbook.worksheets)}."
                )
            worksheet = workbook.worksheets[selected]
        else:
            if selected not in workbook.sheetnames:
                raise ValueError(
                    f"Aba '{selected}' não encontrada. Abas: {workbook.sheetnames}."
                )
            worksheet = workbook[selected]

        rows = worksheet.iter_rows(
            min_row=header_row + 1,
            values_only=True,
        )
        header_values = next(rows, None)
        if header_values is None:
            raise ValueError("O XLSX exportado não possui cabeçalho.")
        header = [
            "" if value is None else str(value).replace("\ufeff", "").strip()
            for value in header_values
        ]
        while header and header[-1] == "":
            header.pop()
        validated_header = tuple(COLUMN_ALIASES.get(name, name) for name in header)
        if validated_header != EXPECTED_COLUMNS:
            raise ValueError(
                "O schema ou a ordem da tabela Procedimentos mudou. "
                f"Esperado: {list(EXPECTED_COLUMNS)}. Encontrado: {header}. "
                f"Aliases aceitos: {COLUMN_ALIASES}."
            )

        data_rows = 0
        for row in rows:
            if any(value not in (None, "") for value in row[len(EXPECTED_COLUMNS) :]):
                raise ValueError(
                    "A exportação possui dados em colunas sem cabeçalho após "
                    f"'{EXPECTED_COLUMNS[-1]}'."
                )
            if any(
                value not in (None, "")
                for value in row[: len(EXPECTED_COLUMNS)]
            ):
                data_rows += 1
        if data_rows == 0 and not allow_empty:
            raise ValueError(
                "A exportação da tabela Procedimentos está vazia; o arquivo anterior "
                "não será substituído. Use --allow-empty somente se isso for esperado."
            )
        return WorkbookValidation(
            sheet_name=worksheet.title,
            data_rows=data_rows,
            file_size=path.stat().st_size,
        )
    finally:
        workbook.close()


def _temporary_sibling(path: Path) -> Path:
    descriptor, name = tempfile.mkstemp(
        prefix=f".{path.stem}.",
        suffix=".xlsx.tmp",
        dir=path.parent,
    )
    os.close(descriptor)
    return Path(name)


def _fsync(path: Path) -> None:
    with open(path, "r+b") as file:
        file.flush()
        os.fsync(file.fileno())


def publish_procedures_workbook(
    *,
    source: Path,
    procedures_dir: Path,
    sheet: str,
    header_row: int,
    allow_empty: bool,
    lock_timeout: float,
    delete_source: bool,
) -> tuple[Path, WorkbookValidation, bool]:
    validation = validate_procedures_workbook(
        source,
        sheet=sheet,
        header_row=header_row,
        allow_empty=allow_empty,
    )
    procedures_dir.mkdir(parents=True, exist_ok=True)
    target = procedures_dir / FINAL_FILENAME
    if source.resolve() == target.resolve():
        return target, validation, target.exists()

    with FileLock(Path(f"{target}.lock"), timeout=lock_timeout):
        replaced_existing = target.is_file()
        temporary = _temporary_sibling(target)
        try:
            shutil.copy2(source, temporary)
            _fsync(temporary)
            try:
                os.replace(temporary, target)
            except PermissionError as exc:
                raise PermissionError(
                    f"Não foi possível substituir {target}. Feche Procedimentos.xlsx "
                    "no Excel e tente novamente."
                ) from exc
        finally:
            temporary.unlink(missing_ok=True)

    if delete_source:
        try:
            source.unlink(missing_ok=True)
        except OSError as exc:
            LOGGER.warning(
                "Procedimentos.xlsx foi publicado, mas o download temporário não pôde "
                "ser removido: %s",
                exc,
            )
    return target, validation, replaced_existing


def _is_pipefy_url(url: str) -> bool:
    host = (urlparse(url).hostname or "").lower()
    return (
        host == "pipefy.com"
        or host.endswith(".pipefy.com")
        or host == "staticpipefy.com"
        or host.endswith(".staticpipefy.com")
    )


def _redact(headers: dict[str, Any], enabled: bool) -> dict[str, str]:
    return {
        str(key): (
            "<REDACTED>"
            if enabled and str(key).lower() in SENSITIVE_HEADERS
            else str(value)
        )
        for key, value in (headers or {}).items()
    }


def _json_object(value: str | None) -> Any:
    if not value:
        return None
    try:
        return json.loads(value)
    except (TypeError, json.JSONDecodeError):
        return None


def save_network_capture(driver: Any, output: Path, redact: bool) -> int:
    requests_by_id: dict[str, dict[str, Any]] = {}
    responses_by_id: dict[str, dict[str, Any]] = {}
    request_extra: dict[str, dict[str, Any]] = {}
    response_extra: dict[str, dict[str, Any]] = {}
    order: list[str] = []

    for entry in driver.get_log("performance"):
        try:
            message = json.loads(entry["message"])["message"]
        except (KeyError, TypeError, json.JSONDecodeError):
            continue
        method = message.get("method")
        params = message.get("params", {})
        request_id = params.get("requestId")
        if not request_id:
            continue

        if method == "Network.requestWillBeSent":
            request = params.get("request", {})
            if _is_pipefy_url(request.get("url", "")):
                if request_id not in requests_by_id:
                    order.append(request_id)
                requests_by_id[request_id] = request
        elif method == "Network.requestWillBeSentExtraInfo":
            request_extra[request_id] = params.get("headers", {})
        elif method == "Network.responseReceived":
            response = params.get("response", {})
            if _is_pipefy_url(response.get("url", "")):
                responses_by_id[request_id] = response
        elif method == "Network.responseReceivedExtraInfo":
            response_extra[request_id] = params.get("headers", {})

    records: list[dict[str, Any]] = []
    for request_id in order:
        request = requests_by_id[request_id]
        response = responses_by_id.get(request_id, {})
        body = request.get("postData")
        payload = _json_object(body)
        operation_name = payload.get("operationName") if isinstance(payload, dict) else None
        query = payload.get("query") if isinstance(payload, dict) else None
        variables = payload.get("variables") if isinstance(payload, dict) else None
        searchable = " ".join(
            str(value)
            for value in (request.get("url"), operation_name, query, body)
            if value
        ).lower()
        classification = (
            "graphql_export"
            if "export" in searchable and ("graphql" in searchable or query)
            else "graphql"
            if "graphql" in searchable or query
            else "http_pipefy"
        )
        records.append(
            {
                "index": len(records),
                "captured_at": now_iso(),
                "request_id": request_id,
                "method": request.get("method", ""),
                "url": request.get("url", ""),
                "status_code": response.get("status"),
                "classification": classification,
                "request_headers": _redact(
                    {
                        **(request.get("headers") or {}),
                        **request_extra.get(request_id, {}),
                    },
                    redact,
                ),
                "response_headers": _redact(
                    {
                        **(response.get("headers") or {}),
                        **response_extra.get(request_id, {}),
                    },
                    redact,
                ),
                "request_body": body,
                "graphql_operation_name": operation_name,
                "graphql_query": query,
                "graphql_variables": variables,
            }
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    with open(output, "w", encoding="utf-8") as file:
        for record in records:
            file.write(json.dumps(record, ensure_ascii=False) + "\n")
    return len(records)


def _browser_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--email", default=DEFAULT_EMAIL)
    parser.add_argument("--database-url", default=DEFAULT_DATABASE_URL)
    parser.add_argument("--user-data-dir", default=DEFAULT_EDGE_PROFILE)
    parser.add_argument("--step-timeout", type=int, default=60)
    parser.add_argument("--sso-timeout", type=int, default=300)
    parser.add_argument(
        "--browser-start-timeout",
        type=int,
        default=900,
        help="Tempo máximo aguardando o perfil Edge ficar livre (s).",
    )


def _output_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument(
        "--procedures-dir",
        default=DEFAULT_PROCEDURES_DIR,
        help="Pasta final; o arquivo será sempre Procedimentos.xlsx.",
    )
    parser.add_argument("--sheet", default="0", help="Nome ou índice da aba.")
    parser.add_argument("--header-row", type=int, default=0)
    parser.add_argument("--allow-empty", action="store_true")
    parser.add_argument("--lock-timeout", type=float, default=30.0)


def _run_headless_check(args: argparse.Namespace) -> None:
    driver, _ = _open_database_with_login_fallback(
        headless=True,
        email=args.email,
        database_url=args.database_url,
        user_data_dir=args.user_data_dir,
        step_timeout=args.step_timeout,
        sso_timeout=args.sso_timeout,
        browser_start_timeout=args.browser_start_timeout,
    )
    driver.quit()


def command_login(args: argparse.Namespace) -> None:
    _perform_visible_login(
        email=args.email,
        database_url=args.database_url,
        user_data_dir=args.user_data_dir,
        step_timeout=args.step_timeout,
        sso_timeout=args.sso_timeout,
        browser_start_timeout=args.browser_start_timeout,
    )
    _mark_daily_login(args.user_data_dir)
    LOGGER.info("Perfil autenticado e pronto para execução headless.")


def command_check_headless(args: argparse.Namespace) -> None:
    _run_headless_check(args)
    LOGGER.info("Teste headless aprovado para a tabela Procedimentos.")


def command_run(args: argparse.Namespace) -> None:
    procedures_dir = resolve_path(args.procedures_dir)
    staging_dir = resolve_path(DEFAULT_PROCEDURES_STAGING_DIR)
    target = procedures_dir / FINAL_FILENAME
    capture_path = resolve_path(args.capture_out) if args.capture_out else None

    if args.weekdays_only and datetime.now().astimezone().weekday() >= 5:
        LOGGER.info("Execução ignorada: o agendamento roda somente de segunda a sexta.")
        return

    if args.ensure_daily_login:
        _ensure_daily_login(
            email=args.email,
            database_url=args.database_url,
            user_data_dir=args.user_data_dir,
            step_timeout=args.step_timeout,
            sso_timeout=args.sso_timeout,
            browser_start_timeout=args.browser_start_timeout,
        )

    with FileLock(Path(f"{target}.run.lock"), timeout=args.lock_timeout):
        driver, export_button = _open_database_with_login_fallback(
            headless=args.headless,
            email=args.email,
            database_url=args.database_url,
            user_data_dir=args.user_data_dir,
            step_timeout=args.step_timeout,
            sso_timeout=args.sso_timeout,
            download_dir=staging_dir,
            capture_network=bool(capture_path),
            browser_start_timeout=args.browser_start_timeout,
        )
        downloaded: Path | None = None
        capture_saved = False
        try:
            if capture_path:
                driver.get_log("performance")

            before = directory_snapshot(staging_dir)
            LOGGER.info("Clicando em Export na tabela Procedimentos.")
            _click(driver, export_button)
            downloaded = wait_for_download(
                staging_dir,
                before,
                args.download_timeout,
            )
            LOGGER.info("Download concluído: %s", downloaded)

            if capture_path:
                total = save_network_capture(driver, capture_path, args.redact)
                capture_saved = True
                LOGGER.info(
                    "Captura CDP salva em %s (%s requisições).",
                    capture_path,
                    total,
                )
        finally:
            if capture_path and not capture_saved:
                try:
                    total = save_network_capture(driver, capture_path, args.redact)
                    LOGGER.info(
                        "Captura CDP parcial salva em %s (%s requisições).",
                        capture_path,
                        total,
                    )
                except Exception as exc:
                    LOGGER.warning("Não foi possível salvar a captura CDP: %s", exc)
            driver.quit()

        if downloaded is None:
            raise RuntimeError("O Edge foi encerrado sem produzir o XLSX.")
        final, validation, replaced = publish_procedures_workbook(
            source=downloaded,
            procedures_dir=procedures_dir,
            sheet=args.sheet,
            header_row=args.header_row,
            allow_empty=args.allow_empty,
            lock_timeout=args.lock_timeout,
            delete_source=True,
        )
        LOGGER.info(
            "%s publicado: linhas=%s, aba=%s, tamanho=%s bytes, substituiu_anterior=%s.",
            final,
            validation.data_rows,
            validation.sheet_name,
            validation.file_size,
            replaced,
        )


def command_replace(args: argparse.Namespace) -> None:
    final, validation, replaced = publish_procedures_workbook(
        source=resolve_path(args.input),
        procedures_dir=resolve_path(args.procedures_dir),
        sheet=args.sheet,
        header_row=args.header_row,
        allow_empty=args.allow_empty,
        lock_timeout=args.lock_timeout,
        delete_source=False,
    )
    LOGGER.info(
        "%s publicado sem Selenium: linhas=%s, aba=%s, substituiu_anterior=%s.",
        final,
        validation.data_rows,
        validation.sheet_name,
        replaced,
    )


def _powershell_quote(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def _scheduled_python_executable() -> str:
    """Prefere pythonw.exe para o agendamento não abrir uma janela de console."""

    current = Path(sys.executable).resolve()
    if os.name != "nt":
        return str(current)

    candidates = [current.with_name("pythonw.exe")]
    discovered = shutil.which("pythonw.exe")
    if discovered:
        candidates.append(Path(discovered).resolve())

    for candidate in candidates:
        if candidate.is_file():
            return str(candidate)

    LOGGER.warning(
        "pythonw.exe não foi encontrado; usando %s. O modo --silent continuará "
        "ativo, mas o Windows poderá exibir uma janela de console.",
        current,
    )
    return str(current)


def _run_windows_powershell(script: str) -> str:
    if os.name != "nt":
        raise RuntimeError("O agendamento só pode ser instalado no Windows.")
    executable = shutil.which("powershell.exe")
    if not executable:
        raise FileNotFoundError("powershell.exe não encontrado.")
    encoded = base64.b64encode(script.encode("utf-16le")).decode("ascii")
    result = subprocess.run(
        [
            executable,
            "-NoLogo",
            "-NoProfile",
            "-NonInteractive",
            "-EncodedCommand",
            encoded,
        ],
        capture_output=True,
        text=True,
        check=False,
        creationflags=subprocess.CREATE_NO_WINDOW,
    )
    if result.returncode:
        detail = (result.stderr or result.stdout).strip()
        raise RuntimeError(f"Falha no Agendador de Tarefas: {detail}")
    return result.stdout.strip()


def command_install_schedule(args: argparse.Namespace) -> None:
    LOGGER.info("Validando headless antes de criar o agendamento de Procedimentos.")
    _run_headless_check(args)

    python = _scheduled_python_executable()
    script = str(Path(__file__).resolve())
    procedures_dir = str(resolve_path(args.procedures_dir))
    log_file = str(resolve_path(args.log_file))
    command = [
        script,
        "--log-file",
        log_file,
        "--silent",
        "run",
        "--headless",
        "--weekdays-only",
        "--ensure-daily-login",
        "--email",
        args.email,
        "--database-url",
        args.database_url,
        "--user-data-dir",
        args.user_data_dir,
        "--step-timeout",
        str(args.step_timeout),
        "--sso-timeout",
        str(args.sso_timeout),
        "--browser-start-timeout",
        str(args.browser_start_timeout),
        "--download-timeout",
        str(args.download_timeout),
        "--procedures-dir",
        procedures_dir,
        "--sheet",
        args.sheet,
        "--header-row",
        str(args.header_row),
        "--lock-timeout",
        str(args.lock_timeout),
    ]
    if args.allow_empty:
        command.append("--allow-empty")
    arguments = subprocess.list2cmdline(command)
    task_name = _powershell_quote(args.task_name)

    weekdays = ",".join(TASK_WEEKDAYS)
    # Cada horário vira um gatilho semanal limitado a segunda-feira a sexta-feira.
    trigger_commands = "\n".join(
        "$triggers += New-ScheduledTaskTrigger -Weekly -WeeksInterval 1 "
        f"-DaysOfWeek {weekdays} -At '{task_time}'"
        for task_time in TASK_TIMES
    )
    schedule_description = (
        "Substitui procedimentos\\Procedimentos.xlsx no PBi de segunda a sexta às "
        + ", ".join(TASK_TIMES)
        + "; a primeira execução roda o login visível do Pipefy."
    )

    ps_script = f"""
$ErrorActionPreference = 'Stop'
$action = New-ScheduledTaskAction `
    -Execute {_powershell_quote(python)} `
    -Argument {_powershell_quote(arguments)} `
    -WorkingDirectory {_powershell_quote(str(BASE_DIR))}

$triggers = @()
{trigger_commands}

$settings = New-ScheduledTaskSettingsSet `
    -StartWhenAvailable `
    -MultipleInstances IgnoreNew `
    -ExecutionTimeLimit (New-TimeSpan -Hours 2)
$user = [System.Security.Principal.WindowsIdentity]::GetCurrent().Name
$principal = New-ScheduledTaskPrincipal `
    -UserId $user `
    -LogonType Interactive `
    -RunLevel Limited
Register-ScheduledTask `
    -TaskName {task_name} `
    -Action $action `
    -Trigger $triggers `
    -Settings $settings `
    -Principal $principal `
    -Description {_powershell_quote(schedule_description)} `
    -Force | Out-Null
"""
    _run_windows_powershell(ps_script)
    LOGGER.info(
        "Tarefa '%s' instalada de segunda a sexta para %s.",
        args.task_name,
        ", ".join(TASK_TIMES),
    )


def command_uninstall_schedule(args: argparse.Namespace) -> None:
    task_name = _powershell_quote(args.task_name)
    ps_script = f"""
$ErrorActionPreference = 'Stop'
if (Get-ScheduledTask -TaskName {task_name} -ErrorAction SilentlyContinue) {{
    Unregister-ScheduledTask -TaskName {task_name} -Confirm:$false
}}
"""
    _run_windows_powershell(ps_script)
    LOGGER.info("Tarefa '%s' removida.", args.task_name)


def load_capture(path: Path) -> list[dict[str, Any]]:
    with open(path, "r", encoding="utf-8") as file:
        return [json.loads(line) for line in file if line.strip()]


def command_inspect(args: argparse.Namespace) -> None:
    rows = load_capture(resolve_path(args.input))
    selected = [
        row
        for row in rows
        if not args.only_graphql
        or str(row.get("classification", "")).startswith("graphql")
    ]
    LOGGER.info("Requisições lidas=%s, selecionadas=%s.", len(rows), len(selected))
    for row in selected:
        print(
            f"[{row.get('index')}] {row.get('method')} {row.get('status_code')} "
            f"{row.get('classification')} "
            f"op={row.get('graphql_operation_name') or '-'}\n"
            f"    {row.get('url')}"
        )


def _select_capture(
    rows: Sequence[dict[str, Any]],
    index: int,
) -> dict[str, Any]:
    for row in rows:
        if int(row.get("index", -1)) == index:
            return row
    raise IndexError(f"Índice de captura não encontrado: {index}")


def command_requests(args: argparse.Namespace) -> None:
    row = _select_capture(load_capture(resolve_path(args.input)), args.index)
    excluded = {"content-length", "host", "accept-encoding", "priority"}
    headers = {
        key: value
        for key, value in (row.get("request_headers") or {}).items()
        if value != "<REDACTED>"
        and not key.startswith(":")
        and key.lower() not in excluded
        and not key.lower().startswith(("sec-fetch-", "sec-ch-ua"))
    }
    body = row.get("request_body")
    payload = _json_object(body)
    output = resolve_path(args.out)
    output.parent.mkdir(parents=True, exist_ok=True)
    script = f"""#!/usr/bin/env python3
from pathlib import Path
import requests

url = {row.get("url", "")!r}
method = {row.get("method", "POST")!r}
headers = {headers!r}
json_payload = {payload!r}
raw_payload = {None if payload is not None else body!r}

response = requests.request(
    method,
    url,
    headers=headers,
    json=json_payload,
    data=raw_payload,
    timeout=120,
)
response.raise_for_status()

disposition = response.headers.get("content-disposition", "").lower()
content_type = response.headers.get("content-type", "").lower()
if "attachment" in disposition or "spreadsheet" in content_type:
    target = Path({str(Path(DEFAULT_PROCEDURES_DIR) / FINAL_FILENAME)!r})
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(response.content)
    print(f"Arquivo salvo em: {{target.resolve()}}")
else:
    print(response.status_code)
    print(response.text[:4000])
"""
    output.write_text(script, encoding="utf-8")
    LOGGER.info("Script requests salvo em %s.", output)
    if any(
        value == "<REDACTED>"
        for value in (row.get("request_headers") or {}).values()
    ):
        LOGGER.warning(
            "A captura está mascarada; recapture com --no-redact para reproduzir "
            "headers de autenticação."
        )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Edge e msedgedriver v150 ou superior: exporta a tabela "
            "Procedimentos do Pipefy e substitui procedimentos/Procedimentos.xlsx no PBi."
        )
    )
    parser.add_argument("--verbose", action="store_true")
    parser.add_argument(
        "--silent",
        action="store_true",
        help="Não escreve no console; mantém somente o arquivo de log.",
    )
    parser.add_argument("--log-file", default=DEFAULT_LOG_FILE)
    subparsers = parser.add_subparsers(dest="command", required=True)

    login = subparsers.add_parser("login", help="Autentica com janela visível.")
    _browser_arguments(login)
    login.set_defaults(func=command_login)

    check = subparsers.add_parser(
        "check-headless",
        help="Valida a sessão sem exportar.",
    )
    _browser_arguments(check)
    check.set_defaults(func=command_check_headless)

    run = subparsers.add_parser(
        "run",
        aliases=["flow"],
        help="Exporta e substitui procedimentos/Procedimentos.xlsx no projeto PBi.",
    )
    _browser_arguments(run)
    _output_arguments(run)
    run.add_argument("--headless", action="store_true")
    run.add_argument(
        "--weekdays-only",
        action="store_true",
        help="Ignora a execução aos sábados e domingos.",
    )
    run.add_argument(
        "--ensure-daily-login",
        action="store_true",
        help="Executa o login visível compartilhado na primeira rodada do dia.",
    )
    run.add_argument("--download-timeout", type=int, default=300)
    run.add_argument("--capture-out", help="JSONL opcional da chamada de export.")
    run.set_defaults(redact=True)
    run.add_argument(
        "--no-redact",
        dest="redact",
        action="store_false",
        help="Mantém tokens/cookies na captura local; use com cautela.",
    )
    run.set_defaults(func=command_run)

    replace = subparsers.add_parser(
        "replace",
        help="Valida e publica um XLSX já baixado sem Selenium.",
    )
    replace.add_argument("--input", required=True)
    _output_arguments(replace)
    replace.set_defaults(func=command_replace)

    schedule = subparsers.add_parser(
        "install-schedule",
        help=(
            "Valida headless e agenda de segunda a sexta às "
            "09:10, 11:10, 13:10, 15:10 e 17:10."
        ),
    )
    _browser_arguments(schedule)
    _output_arguments(schedule)
    schedule.add_argument("--download-timeout", type=int, default=300)
    schedule.add_argument("--task-name", default=DEFAULT_TASK_NAME)
    schedule.set_defaults(func=command_install_schedule)

    unschedule = subparsers.add_parser(
        "uninstall-schedule",
        help="Remove a tarefa do Agendador do Windows.",
    )
    unschedule.add_argument("--task-name", default=DEFAULT_TASK_NAME)
    unschedule.set_defaults(func=command_uninstall_schedule)

    inspect = subparsers.add_parser(
        "inspect",
        help="Lista as requisições de uma captura CDP.",
    )
    inspect.add_argument("--input", required=True)
    inspect.add_argument("--only-graphql", action="store_true")
    inspect.set_defaults(func=command_inspect)

    requests_parser = subparsers.add_parser(
        "requests",
        help="Converte uma requisição capturada em Python requests.",
    )
    requests_parser.add_argument("--input", required=True)
    requests_parser.add_argument("--index", type=int, required=True)
    requests_parser.add_argument("--out", required=True)
    requests_parser.set_defaults(func=command_requests)
    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()
    configure_logging(args.verbose, args.log_file, args.silent)
    try:
        args.func(args)
    except KeyboardInterrupt:
        LOGGER.error("Execução cancelada.")
        raise SystemExit(130)
    except Exception as exc:
        if args.verbose:
            LOGGER.exception("Falha: %s", exc)
        else:
            LOGGER.error("Falha: %s", exc)
        raise SystemExit(1)


if __name__ == "__main__":
    main()
