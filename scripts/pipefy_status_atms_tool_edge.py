#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""Exporta Status Projeto ATMS e substitui o XLSX no compartilhamento.

Instalação: copiar este arquivo para extraction\scripts. O orchestrator 1.2.4
já descobre pipefy_*_tool_edge.py e executa --log-file <log> run --headless
--weekdays-only. Não exige alterações no JSON ou novas tarefas agendadas.

Fluxo: organização 300470324 -> relatório carregado -> exportação -> download.
Perfil e usuário compartilhados com os outros extratores.
Origem: pipe 307055244, report 301058637.
URL: https://app.pipefy.com/pipes/307055244/reports_v2/301058637.
Botão de exportação: button[data-testid="report-export"] (Export Report).
Download: [power_bi]_status_projeto_atms_dd-MM-yyyy.xlsx.
Destino: \\10.69.17.37\01.Files_NaoApagar\Report_Pipefy\status_projeto_atms.xlsx.

Dependências: selenium>=4.34.2,<5; openpyxl>=3.1,<4 (já no requirements.txt).
SSO/perfil/driver seguem o pacote 1.2.4; nenhum token de API é necessário.

O destino ATMS é próprio: PIPEFY_OUTPUT_DIR do orchestrator não o redireciona.
Use --output-dir ou PIPEFY_ATMS_OUTPUT_DIR para uma alteração explícita.
A ordem das 57 colunas é validada. Entidades HTML, espaços e Unicode são
normalizados APENAS para comparação; os bytes do XLSX são preservados.
O arquivo anterior nunca é apagado antes: a nova cópia é validada na pasta
final e publicada com os.replace. Falhas de validação/cópia/substituição
preservam o anterior. Exportações vazias exigem --allow-empty.

A exportação acompanha as respostas emitidas pela própria página. Se ela
entregar o link do XLSX sem iniciar o download, o Edge abre esse link na mesma
sessão. Não é feita uma exportação paralela por API. Falhas geram screenshot
atms_failure.png e diagnóstico atms_diagnostic.json na pasta da execução,
sem gravar cookies, tokens, cabeçalhos HTTP ou URLs assinadas.

Exemplos (PowerShell, a partir de extraction):
  py .\scripts\pipefy_status_atms_tool_edge.py run --headless
  py .\scripts\pipefy_status_atms_tool_edge.py login
  py .\scripts\pipefy_status_atms_tool_edge.py replace --input "C:\Downloads\[power_bi]_status_projeto_atms_29-09-2026.xlsx"
  py .\pipefy_orchestrator.py --config .\pipefy_orchestrator.config.json run --job pipefy_status_atms_tool_edge --delay-seconds 0

Rota oficial de relatórios:
https://developers.pipefy.com/reference/get-reports
Exportação pela interface:
https://help.pipefy.com/en/articles/619297-how-to-create-pipe-reports
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import html
import json
import logging
import os
import re
import shutil
import subprocess
import tempfile
import time
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from logging.handlers import RotatingFileHandler
from pathlib import Path, PureWindowsPath
from typing import Any, Sequence
from urllib.parse import urlparse

APP_NAME = "pipefy_status_atms_tool"
SCRIPT_VERSION = "1.1.0"
BASE_DIR = Path(__file__).resolve().parent
EXTRACTION_DIR = BASE_DIR.parent
DEFAULT_OUTPUT_DIR = os.environ.get(
    "PIPEFY_ATMS_OUTPUT_DIR", r"\\10.69.17.37\01.Files_NaoApagar\Report_Pipefy"
)
DEFAULT_WORK_DIR = Path(os.environ.get(
    "PIPEFY_WORK_DIR", str(EXTRACTION_DIR / ".pipefy_orchestrator" / "downloads")
)).expanduser()
DEFAULT_DOWNLOAD_DIR = str(DEFAULT_WORK_DIR / "status_atms")
DEFAULT_LOG_FILE = str(EXTRACTION_DIR / "logs" / "pipefy_status_atms_tool.log")
EDGE_DRIVER_PATH = Path(os.environ.get(
    "PIPEFY_EDGE_DRIVER_PATH",
    r"C:\Users\chailer\Downloads\edgedriver_win64\msedgedriver.exe",
))
MIN_EDGE_MAJOR_VERSION = 150
DEFAULT_EDGE_PROFILE = r"C:\SeleniumEdgeProfile"
DEFAULT_EMAIL = "chailer@kpmg.com.br"
ORGANIZATION_ID = "300470324"
DEFAULT_ORGANIZATION_URL = f"https://app.pipefy.com/organizations/{ORGANIZATION_ID}"
PIPE_ID = "307055244"
REPORT_ID = "301058637"
DEFAULT_REPORT_URL = f"https://app.pipefy.com/pipes/{PIPE_ID}/reports_v2/{REPORT_ID}"
FINAL_FILENAME = "status_projeto_atms.xlsx"
DOWNLOAD_NAME_RE = re.compile(
    r"^\[power_bi\]_status_projeto_atms_(\d{2}-\d{2}-\d{4})(?: \(\d+\))?\.xlsx$",
    re.IGNORECASE,
)

EXPECTED_COLUMNS = (
    "Current phase",
    "Assignees",
    "Nome do Cliente",
    "Financial/Non-Financial",
    "Qual o segmento do cliente?",
    "Data de abertura da Solicitação",
    "Data Captação",
    "Data Início Levantamento Requisitos",
    "Data Fim Levantamento de Requisitos (Expectativa)",
    "Data Fim Levantamento de Requisitos (Real)",
    "Nome da Rotina",
    "Status da Solicitação",
    "AUDIT - Valor Manual por Ano/Eficiência",
    "AUDIT - Horas Manuais por Ano/Eficiência (Assistente)",
    "ATMS - Horas Totais",
    "ATMS - Valor Total",
    "Data de Início Desenvolvimento",
    "Data Prevista do Final do Desenvolvimento",
    "Data Final Desenvolvimento",
    "Status do Desenvolvimento",
    "Data de Início Homologação",
    "Data Prevista do Final da Homologação",
    "Data Final Homologação",
    "Data do Início Documentação",
    "Data Prevista do Final da Documentação",
    "Data Final Documentação",
    "Status da Homologação",
    "Data Go Live",
    "Labels",
    "Data de Aprovação Eficiência (Expectativa)",
    "Data da Aprovação Eficiência (Real)",
    "Eficiência Aprovada?",
    "Data de entrada em Planejamento",
    "Data Fim de Planejamento (Expectativa)",
    "Data Planejada Início do Desenvolvimento",
    "Status do Planejamento",
    "Data de Entrada no Backlog",
    "Status do Backlog",
    "Data Prevista Aprovação da Homologação",
    "Homologação Aprovado em",
    "Data que entrou na Fase",
    "Data Prevista para Publicação",
    "O Deploy foi feito?",
    "First time enter Aprovação de Eficiência",
    "Último Fase da Solicitação",
    "ID",
    "Houve Substituição da Rotina?",
    "Data Planejada Fim Documentação",
    "Data de Retorno de Auditoria",
    "Data Planejada Fim Levantamento de Requisitos",
    "Data Planejada do Final do Desenvolvimento",
    "Data Planejada do Final da Homologação",
    "Data de Entrada em Espera",
    "Qual o engagement do cliente?",
    "Data do Arquivamento",
    "AUDIT - Horas Manuais por Execução (Assistente)",
    "AUDIT - Valor Manual por Execução",
)

# Rótulos das opções de download que podem aparecer após Export Report.
DOWNLOAD_LABELS = (
    "Download", "Baixar", "Download report", "Baixar relatório",
    "Download Excel", "Download Excel file", "Download as Excel",
    "Download as XLSX", "Download XLSX", "Download (.xlsx)",
    "Download report (.xlsx)", "Download report in Excel",
    "Export to Excel", "Exportar para Excel", "Exportar Excel",
    "Baixar Excel", "Baixar XLSX", "Baixar (.xlsx)",
    "Baixar relatório (.xlsx)", "Baixar relatório em Excel", "Excel", ".xlsx",
)


def label_locators(labels: Sequence[str]) -> tuple[tuple[str, str], ...]:
    """Elementos clicáveis com rótulo exato, inclusive menus e ícones."""
    lower = "translate(normalize-space(%s), 'ABCDEFGHIJKLMNOPQRSTUVWXYZ', 'abcdefghijklmnopqrstuvwxyz')"
    tests = []
    for label in labels:
        for field in (".", "@aria-label", "@title", "@data-tooltip"):
            tests.append(f"{lower % field}='{label.lower()}'")
    return (("xpath", "//*[self::button or self::a or @role='button' or @role='menuitem']["
             + " or ".join(tests) + "]"),)


# Botão confirmado no HTML do relatório. As classes do wrapper/SVG são
# geradas pela interface; o clique deve ocorrer no próprio button.
EXPORT_LOCATORS = (
    ("css selector", 'button[data-testid="report-export"]'),
    ("css selector", 'button[aria-label="Export Report"]'),
    ("css selector", 'button[aria-label="Exportar relatório"]'),
)
DOWNLOAD_LOCATORS = label_locators(DOWNLOAD_LABELS)
SSO_SELECTOR = (
    "a[id*='social-KPMGOficial'][id*='saml'], "
    "a[href*='KPMGOficial'][href*='saml']"
)
LOGGER = logging.getLogger(APP_NAME)


def resolve_path(value: str | Path) -> Path:
    raw = os.path.expandvars(str(value))
    if os.name != "nt" and PureWindowsPath(raw).is_absolute():
        raise OSError("Caminhos Windows/UNC exigem Windows. Use --output-dir local para testes.")
    path = Path(raw).expanduser()
    return path.resolve() if path.is_absolute() else (BASE_DIR / path).resolve()


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def configure_logging(verbose: bool, log_file: str, silent: bool = False) -> None:
    LOGGER.handlers.clear()
    LOGGER.propagate = False
    LOGGER.setLevel(logging.DEBUG if verbose else logging.INFO)
    formatter = logging.Formatter(
        "%(asctime)s | %(levelname)s | %(message)s",
        datefmt="%H:%M:%S",
    )

    if not silent:
        console = logging.StreamHandler()
        console.setFormatter(formatter)
        LOGGER.addHandler(console)

    path = resolve_path(log_file)
    path.parent.mkdir(parents=True, exist_ok=True)
    file_handler = RotatingFileHandler(
        path,
        maxBytes=5_000_000,
        backupCount=5,
        encoding="utf-8",
    )
    file_handler.setFormatter(formatter)
    LOGGER.addHandler(file_handler)


class FileLock:
    """Lock simples entre processos, compatível com Windows e Unix."""

    def __init__(self, path: Path, timeout: float = 30.0) -> None:
        self.path = path
        self.timeout = timeout
        self.handle: Any = None

    def __enter__(self) -> "FileLock":
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.handle = open(self.path, "a+b")
        self.handle.seek(0, os.SEEK_END)
        if self.handle.tell() == 0:
            self.handle.write(b"0")
            self.handle.flush()

        deadline = time.monotonic() + self.timeout
        while True:
            try:
                self.handle.seek(0)
                if os.name == "nt":
                    import msvcrt

                    msvcrt.locking(self.handle.fileno(), msvcrt.LK_NBLCK, 1)
                else:
                    import fcntl

                    fcntl.flock(self.handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
                return self
            except (OSError, BlockingIOError):
                if time.monotonic() >= deadline:
                    self.handle.close()
                    self.handle = None
                    raise TimeoutError(
                        f"Outra execução está usando o recurso: {self.path}"
                    )
                time.sleep(0.25)

    def __exit__(self, exc_type: Any, exc: Any, traceback: Any) -> None:
        if not self.handle:
            return
        try:
            self.handle.seek(0)
            if os.name == "nt":
                import msvcrt

                msvcrt.locking(self.handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl

                fcntl.flock(self.handle.fileno(), fcntl.LOCK_UN)
        finally:
            self.handle.close()
            self.handle = None


class ProfileInUseError(RuntimeError):
    pass


def _is_valid_profile_json(data: bytes) -> bool:
    try:
        value = json.loads(
            data.decode("utf-8-sig"),
            parse_constant=lambda token: (_ for _ in ()).throw(
                ValueError(f"constante JSON inválida: {token}")
            ),
        )
    except (UnicodeError, ValueError):
        return False
    return isinstance(value, dict)


def _invalid_profile_backup(path: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    return path.with_name(f"{path.name}.invalid-{timestamp}.bak")


def repair_invalid_edge_profile_json(profile: Path) -> tuple[Path, ...]:
    """Isola Preferences/Local State inválidos sem apagar o perfil ou cookies."""
    repaired: list[Path] = []
    candidates = (
        profile / "Default" / "Preferences",
        profile / "Local State",
    )
    for path in candidates:
        if not path.is_file():
            continue
        try:
            first = path.read_bytes()
        except OSError as exc:
            raise ProfileInUseError(
                f"Não foi possível ler o arquivo de perfil {path}: {exc}"
            ) from exc
        if _is_valid_profile_json(first):
            continue

        time.sleep(0.25)
        try:
            second = path.read_bytes()
        except FileNotFoundError:
            continue
        except OSError as exc:
            raise ProfileInUseError(
                f"Não foi possível reler o arquivo de perfil {path}: {exc}"
            ) from exc
        if _is_valid_profile_json(second):
            continue
        if first != second:
            raise ProfileInUseError(
                f"O perfil está sendo alterado por outro processo: {path}"
            )

        backup = _invalid_profile_backup(path)
        try:
            os.replace(path, backup)
        except OSError as exc:
            raise ProfileInUseError(
                f"Não foi possível isolar o JSON inválido {path}: {exc}"
            ) from exc
        repaired.append(backup)
        LOGGER.warning(
            "JSON inválido do perfil isolado: %s. Backup: %s",
            path,
            backup,
        )
    return tuple(repaired)


class PipefySessionExpiredError(RuntimeError):
    """Indica que uma execução headless encontrou a tela de login."""


def _command_output(command: Sequence[str]) -> str:
    creation_flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
    result = subprocess.run(
        list(command),
        capture_output=True,
        text=True,
        timeout=20,
        check=False,
        creationflags=creation_flags,
    )
    return "\n".join(part for part in (result.stdout, result.stderr) if part).strip()


def _major_version(version: str, component: str) -> int:
    """Extrai a versão principal de um componente Edge."""

    match = re.match(r"\s*(\d+)", version)
    if not match:
        raise RuntimeError(
            f"Não foi possível identificar a versão principal de {component}: "
            f"{version or '<vazia>'}."
        )
    return int(match.group(1))


def _validate_minimum_edge_version(version: str, component: str) -> int:
    """Aceita qualquer versão do componente a partir da mínima configurada."""

    major = _major_version(version, component)
    if major < MIN_EDGE_MAJOR_VERSION:
        raise RuntimeError(
            f"{component} incompatível: é necessária a versão "
            f"{MIN_EDGE_MAJOR_VERSION} ou superior; encontrada {version}."
        )
    return major


def validate_edge_driver() -> str:
    if not EDGE_DRIVER_PATH.is_file():
        raise FileNotFoundError(
            f"msedgedriver v{MIN_EDGE_MAJOR_VERSION} ou superior não encontrado "
            "no caminho configurado: "
            f"{EDGE_DRIVER_PATH}"
        )
    output = _command_output([str(EDGE_DRIVER_PATH), "--version"])
    match = re.search(r"\b(\d+\.\d+\.\d+\.\d+)\b", output)
    if not match:
        raise RuntimeError(
            f"Não foi possível identificar a versão de {EDGE_DRIVER_PATH}. "
            f"Saída: {output or '<vazia>'}"
        )
    version = match.group(1)
    _validate_minimum_edge_version(version, "msedgedriver")
    return version


def _build_edge_once(
    *,
    headless: bool,
    user_data_dir: str,
    download_dir: Path | None,
    capture_network: bool,
    start_minimized: bool,
) -> Any:
    from selenium import webdriver
    from selenium.webdriver.edge.options import Options
    from selenium.webdriver.edge.service import Service

    driver_version = validate_edge_driver()
    profile = Path(user_data_dir).expanduser().resolve()
    profile.mkdir(parents=True, exist_ok=True)
    repaired_profile_files = repair_invalid_edge_profile_json(profile)
    if repaired_profile_files:
        LOGGER.info(
            "Perfil do Edge reparado; %s arquivo(s) inválido(s) isolado(s).",
            len(repaired_profile_files),
        )

    options = Options()
    options.add_argument(f"--user-data-dir={profile}")
    options.add_argument("--no-first-run")
    options.add_argument("--no-default-browser-check")
    options.add_argument("--disable-notifications")
    options.add_argument("--disable-background-mode")
    options.add_argument("--log-level=3")
    if headless:
        options.add_argument("--headless=new")
        options.add_argument("--window-size=1920,1080")
    elif start_minimized:
        options.add_argument("--start-minimized")
    else:
        options.add_argument("--start-maximized")

    prefs: dict[str, Any] = {
        "credentials_enable_service": False,
        "profile.password_manager_enabled": False,
    }
    if download_dir:
        download_dir.mkdir(parents=True, exist_ok=True)
        prefs.update(
            {
                "download.default_directory": str(download_dir),
                "download.prompt_for_download": False,
                "download.directory_upgrade": True,
                "safebrowsing.enabled": True,
            }
        )
    options.add_experimental_option("prefs", prefs)
    if capture_network:
        options.set_capability("ms:loggingPrefs", {"performance": "ALL"})

    LOGGER.info(
        "Iniciando Edge com msedgedriver %s (versão mínima aceita: v%s).",
        driver_version,
        MIN_EDGE_MAJOR_VERSION,
    )
    try:
        driver = webdriver.Edge(
            service=Service(str(EDGE_DRIVER_PATH), log_output=os.devnull),
            options=options,
        )
    except Exception as exc:
        detail = str(exc)
        if (
            "DevToolsActivePort" in detail
            or "user data directory is already in use" in detail.lower()
        ):
            raise ProfileInUseError(
                f"O perfil dedicado está em uso: {profile}"
            ) from exc
        raise RuntimeError(f"Falha ao iniciar Microsoft Edge: {detail}") from exc

    try:
        browser_version = str(driver.capabilities.get("browserVersion", ""))
        browser_major = _validate_minimum_edge_version(
            browser_version,
            "Microsoft Edge",
        )
        driver_major = _major_version(driver_version, "msedgedriver")

        if browser_major != driver_major:
            LOGGER.warning(
                "Versões principais diferentes; a execução continuará: "
                "Edge=%s, msedgedriver=%s. Se o próprio Selenium recusar a "
                "sessão, atualize o executável msedgedriver.",
                browser_version,
                driver_version,
            )
        else:
            LOGGER.info(
                "Versões aceitas: Edge=%s, msedgedriver=%s.",
                browser_version,
                driver_version,
            )

        if start_minimized:
            try:
                driver.minimize_window()
            except Exception as exc:
                LOGGER.warning("Não foi possível minimizar o Edge: %s", exc)

        if download_dir:
            params = {
                "behavior": "allow",
                "downloadPath": str(download_dir),
                "eventsEnabled": True,
            }
            try:
                driver.execute_cdp_cmd("Browser.setDownloadBehavior", params)
            except Exception:
                driver.execute_cdp_cmd("Page.setDownloadBehavior", params)

        if capture_network:
            driver.execute_cdp_cmd(
                "Network.enable",
                {
                    "maxTotalBufferSize": 100_000_000,
                    "maxResourceBufferSize": 10_000_000,
                    "maxPostDataSize": 10_000_000,
                },
            )
            if "performance" not in driver.log_types:
                raise RuntimeError(
                    "Edge iniciou, mas o log CDP 'performance' não está disponível."
                )
        return driver
    except Exception:
        driver.quit()
        raise


def build_edge(
    *,
    headless: bool,
    user_data_dir: str,
    download_dir: Path | None = None,
    capture_network: bool = False,
    browser_start_timeout: int = 900,
    start_minimized: bool = False,
) -> Any:
    deadline = time.monotonic() + browser_start_timeout
    while True:
        try:
            return _build_edge_once(
                headless=headless,
                user_data_dir=user_data_dir,
                download_dir=download_dir,
                capture_network=capture_network,
                start_minimized=start_minimized,
            )
        except ProfileInUseError:
            remaining = int(deadline - time.monotonic())
            if remaining <= 0:
                raise TimeoutError(
                    "O perfil do Edge permaneceu ocupado. Verifique se outro export "
                    "Pipefy ou uma janela manual ainda está em execução."
                )
            LOGGER.warning(
                "Perfil do Edge ocupado; nova tentativa em 15s (%ss restantes).",
                remaining,
            )
            time.sleep(min(15, remaining))


def _first_element(driver: Any, locators: Sequence[tuple[str, str]]) -> Any | None:
    from selenium.common.exceptions import StaleElementReferenceException

    for by, selector in locators:
        for element in driver.find_elements(by, selector):
            try:
                if element.is_displayed() and element.is_enabled():
                    return element
            except StaleElementReferenceException:
                continue
    return None


def _export_button(driver: Any) -> Any | None:
    if not is_expected_report_url(driver.current_url):
        return None
    return _first_element(driver, EXPORT_LOCATORS)


def is_expected_organization_url(url: str) -> bool:
    parsed = urlparse(url or "")
    return (parsed.scheme == "https" and parsed.hostname == "app.pipefy.com"
            and re.fullmatch(rf"/organizations/{ORGANIZATION_ID}(?:/.*)?", parsed.path) is not None)


def safe_page_url(url: str) -> str:
    """Diagnóstico sem query string, fragmento ou links temporários de download."""
    parsed = urlparse(url or "")
    if parsed.hostname == "app.pipefy.com" and parsed.path.startswith(("/pipes/", "/organizations/")):
        return f"https://app.pipefy.com{parsed.path}"
    return f"{parsed.scheme}://{parsed.hostname or ''}/" if parsed.scheme else "<sem URL>"


def _performance_events(driver: Any) -> list[dict[str, Any]]:
    try:
        entries = driver.get_log("performance")
    except Exception:
        return []
    result = []
    for entry in entries:
        try:
            message = json.loads(entry["message"])["message"]
            if isinstance(message, dict):
                result.append(message)
        except (KeyError, TypeError, ValueError):
            continue
    return result


class PageActivity:
    def __init__(self) -> None:
        self.pending: set[str] = set()
        self.last_activity = time.monotonic()

    def poll(self, driver: Any) -> None:
        for event in _performance_events(driver):
            params = event.get("params", {})
            request_id = params.get("requestId", "")
            method = event.get("method")
            if method == "Network.requestWillBeSent":
                host = urlparse(params.get("request", {}).get("url", "")).hostname or ""
                if params.get("type") in {"XHR", "Fetch"} and (host == "pipefy.com" or host.endswith(".pipefy.com")):
                    self.pending.add(request_id)
                    self.last_activity = time.monotonic()
            elif method in {"Network.loadingFinished", "Network.loadingFailed"} and request_id in self.pending:
                self.pending.discard(request_id)
                self.last_activity = time.monotonic()


PAGE_STATE_JS = r"""
const visible = e => !!(e && e.getClientRects().length && getComputedStyle(e).visibility !== 'hidden');
const busy = Array.from(document.querySelectorAll('[aria-busy="true"], [role="progressbar"]')).some(visible);
return {ready: document.readyState === 'complete', busy,
        content: !!document.body && document.body.innerText.trim().length > 20,
        controls: Array.from(document.querySelectorAll('a,button,[role="navigation"],main')).some(visible)};
"""


def _login_username(driver: Any) -> Any | None:
    return _first_element(driver, (("css selector", "input#username,input[name='username'],input[type='email']"),))


def _organization_login_state(driver: Any) -> tuple[str, Any | None] | bool:
    username = _login_username(driver)
    if username is not None:
        return "login", username
    sso = _first_element(driver, (("css selector", SSO_SELECTOR),))
    if sso is not None:
        return "sso", sso
    if is_expected_organization_url(driver.current_url):
        state = driver.execute_script(PAGE_STATE_JS)
        if state and state.get("ready") and state.get("content") and state.get("controls"):
            return "ready", None
    host = urlparse(driver.current_url or "").hostname or ""
    if host and host != "pipefy.com" and not host.endswith(".pipefy.com"):
        return "external_sso", None
    return False


def _wait_page_loaded(driver: Any, *, report: bool, timeout: int) -> Any | None:
    """Espera o documento, os indicadores de carregamento e as chamadas da SPA."""
    from selenium.common.exceptions import TimeoutException
    from selenium.webdriver.support.ui import WebDriverWait

    activity = PageActivity()
    stable_since: float | None = None

    def ready(current: Any) -> Any:
        nonlocal stable_since
        activity.poll(current)
        if _login_username(current) is not None:
            raise PipefySessionExpiredError("A sessão voltou à tela de login durante o carregamento.")
        correct_url = is_expected_report_url(current.current_url) if report else is_expected_organization_url(current.current_url)
        state = current.execute_script(PAGE_STATE_JS) or {}
        button = _export_button(current) if report and correct_url else None
        content_ready = button is not None if report else state.get("content") and state.get("controls")
        if not (correct_url and state.get("ready") and not state.get("busy") and content_ready and not activity.pending):
            stable_since = None
            return False
        now = time.monotonic()
        if stable_since is None:
            stable_since = now
        if now - stable_since < 2.0 or now - activity.last_activity < 2.0:
            return False
        return button if report else True

    try:
        return WebDriverWait(driver, timeout, poll_frequency=0.5).until(ready)
    except TimeoutException as exc:
        phase = "relatório ATMS" if report else "organização"
        raise RuntimeError(
            f"O carregamento da {phase} não terminou em {timeout}s. "
            f"URL: {safe_page_url(driver.current_url)}; chamadas pendentes={len(activity.pending)}."
        ) from exc


def _click(driver: Any, element: Any) -> None:
    from selenium.common.exceptions import ElementClickInterceptedException, ElementNotInteractableException
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", element)
    try:
        element.click()
    except (ElementClickInterceptedException, ElementNotInteractableException):
        driver.execute_script("arguments[0].click();", element)


def click_report_export(driver: Any, *, wrapper: bool = False) -> None:
    from selenium.webdriver.common.action_chains import ActionChains
    button = _export_button(driver)
    if button is None:
        raise RuntimeError("O botão report-export não está disponível no relatório esperado.")
    element = button
    if wrapper:
        # Ancestral presente no HTML fornecido, usado apenas se o primeiro
        # clique não produziu menu, solicitação de exportação ou download.
        element = button.find_element("xpath", "./ancestor::div[@data-testid='trigger'][1]")
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", element)
    ActionChains(driver).move_to_element(element).pause(0.25).click().perform()


def open_report(
    driver: Any, *, email: str, report_url: str, headless: bool,
    step_timeout: int, sso_timeout: int,
    organization_url: str = DEFAULT_ORGANIZATION_URL,
) -> Any:
    from selenium.common.exceptions import TimeoutException
    from selenium.webdriver.support.ui import WebDriverWait

    if not is_expected_organization_url(organization_url):
        raise ValueError("A página inicial deve ser a organização 300470324.")
    if not is_expected_report_url(report_url):
        raise ValueError("A URL deve apontar para o pipe 307055244 e o report 301058637.")
    LOGGER.info("[1/4] Abrindo a organização %s com o perfil compartilhado.", ORGANIZATION_ID)
    driver.get(organization_url)
    wait = WebDriverWait(driver, step_timeout)
    try:
        state, element = wait.until(_organization_login_state)
    except TimeoutException as exc:
        raise RuntimeError("A organização não carregou nem exibiu o login. URL: "
                           + safe_page_url(driver.current_url)) from exc

    if state != "ready":
        if headless:
            raise PipefySessionExpiredError("A organização requer renovação da sessão SSO compartilhada.")
        if state == "login":
            LOGGER.info("Usando o mesmo usuário dos outros extratores para o SSO.")
            element.clear()
            element.send_keys(email)
            _click(driver, wait.until(lambda d: _first_element(d, (("id", "kc-login"),))))

            def sso_or_redirect(current: Any) -> tuple[str, Any | None] | bool:
                link = _first_element(current, (("css selector", SSO_SELECTOR),))
                if link is not None:
                    return "sso", link
                host = urlparse(current.current_url or "").hostname or ""
                if host == "app.pipefy.com" or (host and host != "pipefy.com" and not host.endswith(".pipefy.com")):
                    return "redirect", None
                return False
            try:
                state, element = wait.until(sso_or_redirect)
            except TimeoutException as exc:
                raise RuntimeError("O SSO KPMGOficial não apareceu após informar o usuário.") from exc
        if state == "sso" and element is not None:
            LOGGER.info("Selecionando o provedor KPMGOficial.")
            _click(driver, element)
        LOGGER.info("Aguardando SSO/MFA na mesma conta por até %ss.", sso_timeout)
        try:
            WebDriverWait(driver, sso_timeout).until(
                lambda d: urlparse(d.current_url or "").hostname == "app.pipefy.com"
            )
        except TimeoutException as exc:
            raise RuntimeError("SSO/MFA não foi concluído no tempo limite.") from exc
        if not is_expected_organization_url(driver.current_url):
            driver.get(organization_url)

    _wait_page_loaded(driver, report=False, timeout=step_timeout)
    LOGGER.info("Organização carregada; sessão compartilhada disponível.")
    LOGGER.info("[2/4] Abrindo o relatório ATMS e aguardando o carregamento completo.")
    driver.get(report_url)
    button = _wait_page_loaded(driver, report=True, timeout=step_timeout)
    LOGGER.info("Relatório carregado e botão report-export disponível.")
    return button


def _perform_visible_login(
    *,
    email: str,
    report_url: str,
    organization_url: str = DEFAULT_ORGANIZATION_URL,
    user_data_dir: str,
    step_timeout: int,
    sso_timeout: int,
    browser_start_timeout: int,
    start_minimized: bool = False,
) -> None:
    """Executa o mesmo fluxo do comando ``login`` em uma janela do Edge."""

    driver = build_edge(
        headless=False,
        user_data_dir=user_data_dir,
        capture_network=True,
        browser_start_timeout=browser_start_timeout,
        start_minimized=start_minimized,
    )
    try:
        open_report(
            driver,
            email=email,
            report_url=report_url,
            organization_url=organization_url,
            headless=False,
            step_timeout=step_timeout,
            sso_timeout=sso_timeout,
        )
    finally:
        driver.quit()


def _open_report_with_login_fallback(
    *,
    headless: bool,
    email: str,
    report_url: str,
    organization_url: str = DEFAULT_ORGANIZATION_URL,
    user_data_dir: str,
    step_timeout: int,
    sso_timeout: int,
    download_dir: Path | None = None,
    capture_network: bool = True,
    browser_start_timeout: int = 900,
) -> tuple[Any, Any]:
    """Abre o relatório e renova o SSO em janela visível quando necessário."""

    driver = build_edge(
        headless=headless,
        user_data_dir=user_data_dir,
        download_dir=download_dir,
        capture_network=capture_network,
        browser_start_timeout=browser_start_timeout,
    )
    try:
        button = open_report(
            driver,
            email=email,
            report_url=report_url,
            organization_url=organization_url,
            headless=headless,
            step_timeout=step_timeout,
            sso_timeout=sso_timeout,
        )
        return driver, button
    except PipefySessionExpiredError:
        driver.quit()
        if not headless:
            raise
    except Exception:
        if download_dir is not None:
            save_export_diagnostics(driver, download_dir, phase="carregamento da organização/relatório")
        driver.quit()
        raise

    LOGGER.warning(
        "Sessão Pipefy expirada. Fallback iniciado: abrindo o Edge visível "
        "para renovar o SSO/MFA."
    )
    _perform_visible_login(
        email=email,
        report_url=report_url,
        organization_url=organization_url,
        user_data_dir=user_data_dir,
        step_timeout=step_timeout,
        sso_timeout=sso_timeout,
        browser_start_timeout=browser_start_timeout,
        start_minimized=True,
    )

    LOGGER.info("Sessão Pipefy renovada. Repetindo a execução headless.")
    retry_driver = build_edge(
        headless=True,
        user_data_dir=user_data_dir,
        download_dir=download_dir,
        capture_network=capture_network,
        browser_start_timeout=browser_start_timeout,
    )
    try:
        button = open_report(
            retry_driver,
            email=email,
            report_url=report_url,
            organization_url=organization_url,
            headless=True,
            step_timeout=step_timeout,
            sso_timeout=sso_timeout,
        )
        return retry_driver, button
    except Exception:
        if download_dir is not None:
            save_export_diagnostics(retry_driver, download_dir, phase="carregamento após renovar SSO")
        retry_driver.quit()
        raise


def _daily_login_paths(user_data_dir: str) -> tuple[Path, Path]:
    """Retorna o estado e o lock compartilhados pelos extratores Pipefy."""

    profile = Path(user_data_dir).expanduser().resolve()
    profile.mkdir(parents=True, exist_ok=True)
    return (
        profile / ".pipefy_daily_login.json",
        profile / ".pipefy_daily_login.lock",
    )


def _daily_login_completed(state_path: Path, today: str) -> bool:
    try:
        state = json.loads(state_path.read_text(encoding="utf-8"))
    except (OSError, TypeError, json.JSONDecodeError):
        return False
    return isinstance(state, dict) and state.get("date") == today


def _write_daily_login_state(state_path: Path, today: str) -> None:
    descriptor, name = tempfile.mkstemp(
        prefix=".pipefy_daily_login.",
        suffix=".json.tmp",
        dir=state_path.parent,
    )
    temporary = Path(name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as file:
            json.dump(
                {
                    "date": today,
                    "authenticated_at": now_iso(),
                    "source": APP_NAME,
                },
                file,
                ensure_ascii=False,
            )
            file.write("\n")
            file.flush()
            os.fsync(file.fileno())
        os.replace(temporary, state_path)
    finally:
        temporary.unlink(missing_ok=True)


def _mark_daily_login(user_data_dir: str) -> None:
    state_path, _ = _daily_login_paths(user_data_dir)
    today = datetime.now().astimezone().date().isoformat()
    _write_daily_login_state(state_path, today)
    LOGGER.info("Login diário do Pipefy confirmado em %s.", today)


def _ensure_daily_login(
    *,
    email: str,
    report_url: str,
    organization_url: str = DEFAULT_ORGANIZATION_URL,
    user_data_dir: str,
    step_timeout: int,
    sso_timeout: int,
    browser_start_timeout: int,
) -> None:
    """Executa uma vez ao dia o fluxo visível equivalente ao comando login."""

    state_path, lock_path = _daily_login_paths(user_data_dir)
    today = datetime.now().astimezone().date().isoformat()
    daily_lock_timeout = max(
        float(browser_start_timeout + sso_timeout + (step_timeout * 2) + 120),
        1800.0,
    )

    with FileLock(lock_path, timeout=daily_lock_timeout):
        if _daily_login_completed(state_path, today):
            LOGGER.info("Login diário do Pipefy já executado em %s.", today)
            return

        LOGGER.info(
            "Primeira execução do dia: executando o fluxo 'login' com janela "
            "visível antes de extrair as bases."
        )
        _perform_visible_login(
            email=email,
            report_url=report_url,
            organization_url=organization_url,
            user_data_dir=user_data_dir,
            step_timeout=step_timeout,
            sso_timeout=sso_timeout,
            browser_start_timeout=browser_start_timeout,
        )
        _write_daily_login_state(state_path, today)
        LOGGER.info("Login diário do Pipefy confirmado em %s.", today)


def is_expected_report_url(url: str) -> bool:
    parsed = urlparse(url or "")
    return (
        parsed.scheme == "https"
        and parsed.hostname == "app.pipefy.com"
        and re.fullmatch(
            rf"/pipes/{PIPE_ID}/(?:reports_v2|reports)/{REPORT_ID}/?", parsed.path
        ) is not None
    )


def is_expected_download_name(name: str) -> bool:
    match = DOWNLOAD_NAME_RE.fullmatch(name)
    if not match:
        return False
    try:
        datetime.strptime(match.group(1), "%d-%m-%Y")
    except ValueError:
        return False
    return True


def directory_snapshot(directory: Path) -> dict[str, tuple[int, int]]:
    snapshot = {}
    for path in directory.iterdir():
        try:
            if path.is_file():
                stat = path.stat()
                snapshot[path.name] = (stat.st_mtime_ns, stat.st_size)
        except FileNotFoundError:
            continue
    return snapshot


def _contains_id(text: str, identifier: str) -> bool:
    return re.search(rf"(?<!\d){re.escape(identifier)}(?!\d)", text) is not None


def _json_nodes(value: Any):
    if isinstance(value, dict):
        yield value
        for item in value.values():
            yield from _json_nodes(item)
    elif isinstance(value, list):
        for item in value:
            yield from _json_nodes(item)


def _safe_error(value: Any) -> str:
    text = re.sub(r"https?://[^\s\"<>]+", "<URL omitida>", str(value))
    return text[:500]


class ExportMonitor:
    """Observa apenas a exportação emitida pela interface, sem criar outra via API."""
    def __init__(self, driver: Any) -> None:
        self.requests: dict[str, dict[str, Any]] = {}
        self.export_ids: set[str] = set()
        self.request_seen = False
        self.download_started = False
        self.download_name: str | None = None
        self.file_url: str | None = None
        self.url_at: float | None = None
        self.url_opened = False
        self.stage = "aguardando efeito do clique"
        self.error: str | None = None
        self.main_handle = driver.current_window_handle
        self.initial_handles = set(driver.window_handles)
        _performance_events(driver)  # Desconsidera chamadas anteriores ao clique.

    def _register_request(self, params: dict[str, Any]) -> None:
        request = params.get("request", {})
        url = request.get("url", "")
        host = urlparse(url).hostname or ""
        if not (host == "pipefy.com" or host.endswith(".pipefy.com")):
            return
        payload = str(request.get("postData", ""))
        searchable = payload + " " + url
        lowered = searchable.lower()
        if "export" in lowered and _contains_id(searchable, REPORT_ID):
            kind = "create"
            self.request_seen = True
            self.stage = "solicitação de exportação enviada pela página"
        elif "pipereportexport" in lowered or any(_contains_id(searchable, i) for i in self.export_ids):
            kind = "status"
        else:
            return
        self.requests[params["requestId"]] = {
            "kind": kind, "payload": searchable, "status": None, "finished": False,
            "attempts": 0, "consumed": False,
        }

    def _consume_body(self, record: dict[str, Any], value: Any) -> None:
        associated = record["kind"] == "create" or any(
            _contains_id(record["payload"], i) for i in self.export_ids
        )
        if not associated:
            return
        for node in _json_nodes(value):
            if node.get("errors"):
                errors = node["errors"]
                messages = [str(e.get("message", "Erro de exportação")) for e in errors if isinstance(e, dict)]
                self.error = "Pipefy recusou a exportação: " + _safe_error("; ".join(messages))
                return
            if record["kind"] == "create":
                export = node.get("pipeReportExport") or node.get("pipe_report_export")
                if isinstance(export, dict) and export.get("id") is not None:
                    self.export_ids.add(str(export["id"]))
                # Formato REST de exportação, quando a resposta já é o objeto.
                if "id" in node and any(k in node for k in ("fileURL", "file_url", "finishedAt")):
                    self.export_ids.add(str(node["id"]))
            if "finishedAt" in node and node["finishedAt"] is None:
                continue
            for key in ("fileURL", "fileUrl", "file_url", "downloadUrl", "download_url"):
                url = node.get(key)
                if not isinstance(url, str):
                    continue
                parsed = urlparse(url)
                if parsed.scheme == "https" and parsed.hostname and not parsed.username and not parsed.password:
                    if self.file_url != url:
                        self.file_url = url
                        self.url_at = time.monotonic()
                        self.stage = "XLSX gerado; link recebido da própria página"
        if self.export_ids and not self.file_url:
            self.stage = "exportação criada; aguardando o XLSX ser gerado"

    def poll(self, driver: Any) -> None:
        for event in _performance_events(driver):
            params = event.get("params", {})
            method = event.get("method", "")
            request_id = params.get("requestId", "")
            if method == "Network.requestWillBeSent":
                request = params.get("request", {})
                host = urlparse(request.get("url", "")).hostname or ""
                if request.get("hasPostData") and not request.get("postData") and (host == "pipefy.com" or host.endswith(".pipefy.com")):
                    try:
                        request["postData"] = driver.execute_cdp_cmd(
                            "Network.getRequestPostData", {"requestId": request_id}
                        ).get("postData", "")
                    except Exception:
                        pass
                self._register_request(params)
            elif request_id in self.requests:
                record = self.requests[request_id]
                if method == "Network.responseReceived":
                    record["status"] = params.get("response", {}).get("status")
                elif method == "Network.loadingFinished":
                    record["finished"] = True
                elif method == "Network.loadingFailed":
                    if record["kind"] == "create":
                        self.error = "Falha na solicitação de exportação: " + _safe_error(params.get("errorText", "rede"))
            elif method in {"Page.downloadWillBegin", "Browser.downloadWillBegin"}:
                self.download_started = True
                self.download_name = params.get("suggestedFilename")
                self.stage = "download iniciado pelo Edge"
            elif method in {"Page.downloadProgress", "Browser.downloadProgress"}:
                if params.get("state") == "canceled":
                    self.error = "O Edge cancelou o download. Confira a política de downloads e as permissões da pasta."

        # Respostas ficam em memória apenas até serem interpretadas; não são
        # salvas em disco. Tentativas curtas toleram o corpo ainda indisponível.
        for request_id, record in list(self.requests.items()):
            if record["consumed"] or not record["finished"]:
                continue
            if record["kind"] != "create" and not any(_contains_id(record["payload"], i) for i in self.export_ids):
                continue
            status = record["status"]
            if status is not None and status >= 400:
                self.error = f"Pipefy retornou HTTP {int(status)} durante a exportação."
                record["consumed"] = True
                continue
            try:
                body = driver.execute_cdp_cmd("Network.getResponseBody", {"requestId": request_id})
                text = body.get("body", "")
                if body.get("base64Encoded"):
                    text = base64.b64decode(text).decode("utf-8")
                self._consume_body(record, json.loads(text))
                record["consumed"] = True
            except (ValueError, TypeError, UnicodeError):
                record["consumed"] = True
            except Exception:
                record["attempts"] += 1
                record["consumed"] = record["attempts"] >= 5
        if self.error:
            raise RuntimeError(self.error)

    def summary(self) -> dict[str, Any]:
        return {"stage": self.stage, "export_request_seen": self.request_seen,
                "export_ids": sorted(self.export_ids), "generated_link_received": bool(self.file_url),
                "generated_link_opened": self.url_opened, "download_started": self.download_started,
                "download_name": self.download_name, "error": self.error}


DOWNLOAD_CONTROLS_JS = r"""
const visible = e => !!(e && e.getClientRects().length && getComputedStyle(e).visibility !== 'hidden');
const norm = s => (s || '').normalize('NFD').replace(/[\u0300-\u036f]/g,'').replace(/\s+/g,' ').trim().toLowerCase();
const seen = new Set(), result = [];
for (const leaf of document.querySelectorAll('button,a,[role="menuitem"],[role="option"],li,span,div')) {
  if (!visible(leaf) || leaf.closest('table,[role="grid"],#username')) continue;
  if (leaf.matches('div,span,li') && leaf.children.length) continue;
  let label = norm(leaf.innerText || leaf.getAttribute('aria-label') || leaf.getAttribute('title'));
  if (!label || label.length > 140 || /email|e-mail|enviar|send|csv|pdf/.test(label)) continue;
  if (!(/^(download|baixar)(\b|\s)/.test(label) || /^(download|baixar|excel|xlsx|\.xlsx)$/.test(label)
        || ((/excel|xlsx/.test(label)) && /export|download|baixar/.test(label)))) continue;
  let e = leaf.closest('button,a,[role="menuitem"],[role="option"],[tabindex]') || leaf;
  if (e.matches('[data-testid="report-export"]') || e.querySelector('[data-testid="report-export"]')) continue;
  if (!visible(e) || e.disabled || e.getAttribute('aria-disabled') === 'true') continue;
  const whole = norm(e.innerText || e.getAttribute('aria-label'));
  if (/email|e-mail|enviar|send|csv|pdf/.test(whole) || whole.length > 160) e = leaf;
  if (seen.has(e)) continue;
  seen.add(e);
  result.push({element:e, label, href:e.getAttribute('href') || ''});
}
return result.slice(0,20);
"""


def _download_controls(driver: Any) -> list[dict[str, Any]]:
    return driver.execute_script(DOWNLOAD_CONTROLS_JS) or []


def _control_key(control: dict[str, Any]) -> tuple[str, str]:
    return str(control.get("label", "")), str(control.get("href", ""))


def _open_generated_download(driver: Any, monitor: ExportMonitor) -> None:
    if not monitor.file_url or monitor.url_opened:
        return
    monitor.url_opened = True
    LOGGER.info("A página gerou o XLSX; abrindo o link recebido na mesma sessão Edge.")
    driver.switch_to.new_window("tab")
    try:
        driver.get(monitor.file_url)
    except Exception as exc:
        # Navegações que se tornam downloads podem ser reportadas como abortadas.
        if "net::ERR_ABORTED" not in str(exc):
            raise RuntimeError("Não foi possível abrir o link do XLSX gerado pela página.") from exc
    finally:
        if monitor.main_handle in driver.window_handles:
            driver.switch_to.window(monitor.main_handle)


def save_export_diagnostics(driver: Any, directory: Path, monitor: ExportMonitor | None = None,
                            *, phase: str = "exportação") -> None:
    directory.mkdir(parents=True, exist_ok=True)
    data: dict[str, Any] = {"version": SCRIPT_VERSION, "at": now_iso(), "phase": phase}
    if monitor is not None:
        data["export"] = monitor.summary()
    try:
        data["page"] = safe_page_url(driver.current_url)
        data["visible_controls"] = driver.execute_script(r"""
          const visible = e => e.getClientRects().length && getComputedStyle(e).visibility !== 'hidden';
          return Array.from(document.querySelectorAll('button,[role="menuitem"],[role="dialog"],[role="menu"],[data-testid="trigger"]'))
            .filter(visible).slice(0,60).map(e => ({tag:e.tagName, role:e.getAttribute('role'),
              testid:e.getAttribute('data-testid'), label:e.getAttribute('aria-label'),
              text:(e.innerText || '').trim().slice(0,180)}));
        """)
        driver.save_screenshot(str(directory / "atms_failure.png"))
    except Exception:
        data["browser_diagnostic_unavailable"] = True
    try:
        (directory / "atms_diagnostic.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        LOGGER.info("Diagnóstico da etapa %s salvo em: %s", phase, directory)
    except OSError:
        LOGGER.warning("Não foi possível salvar o diagnóstico na pasta da execução.")


def wait_for_download(
    directory: Path, before: dict[str, tuple[int, int]], timeout: float, *,
    driver: Any | None = None, monitor: ExportMonitor | None = None,
    baseline_controls: set[tuple[str, str]] | None = None,
) -> Path:
    """Acompanha menu, geração e download; publica somente um XLSX completo."""
    started = time.monotonic()
    deadline = started + timeout
    stable: dict[str, tuple[tuple[int, int], int]] = {}
    clicked_controls: set[tuple[str, str]] = set(baseline_controls or ())
    ui_action = False
    retried_trigger = False
    last_progress = started
    known_handles = set(driver.window_handles) if driver is not None else set()
    while time.monotonic() < deadline:
        if monitor is not None:
            monitor.poll(driver)
        snapshot = directory_snapshot(directory)
        changed = {name: sig for name, sig in snapshot.items() if before.get(name) != sig}
        partials = [name for name in changed if name.lower().endswith((".crdownload", ".tmp"))]
        candidates = [name for name in changed if is_expected_download_name(name)
                      or (monitor is not None and monitor.request_seen
                          and (name == monitor.download_name or bool(monitor.file_url))
                          and name.lower().endswith('.xlsx'))]
        if candidates and not partials:
            if len(candidates) > 1:
                raise RuntimeError("Mais de um XLSX ATMS novo na mesma execução; publicação recusada.")
            name = candidates[0]
            signature = changed[name]
            old_signature, checks = stable.get(name, ((-1, -1), 0))
            checks = checks + 1 if signature == old_signature and signature[1] > 0 else 0
            stable[name] = (signature, checks)
            if checks >= 2:
                return directory / name
        else:
            stable.clear()

        now = time.monotonic()
        if driver is not None and not candidates and not partials:
            # Uma janela nova pode conter a opção de download. Só inspeciona
            # páginas da própria Pipefy, nunca telas de login de outro provedor.
            handles = list(driver.window_handles)
            new_handles = [h for h in handles if h not in known_handles]
            for handle in new_handles:
                try:
                    driver.switch_to.window(handle)
                    if driver.current_url in {"", "about:blank"}:
                        # A aba pode nascer vazia e receber a página depois.
                        # Mantém pendente para inspecioná-la no próximo ciclo.
                        if monitor is not None and monitor.main_handle in driver.window_handles:
                            driver.switch_to.window(monitor.main_handle)
                        continue
                    known_handles.add(handle)
                    host = urlparse(driver.current_url or "").hostname or ""
                    if not (host == "app.pipefy.com") and monitor is not None and monitor.main_handle in driver.window_handles:
                        driver.switch_to.window(monitor.main_handle)
                except Exception:
                    if monitor is not None and monitor.main_handle in driver.window_handles:
                        driver.switch_to.window(monitor.main_handle)
            host = urlparse(driver.current_url or "").hostname or ""
            if host == "app.pipefy.com":
                for control in _download_controls(driver):
                    key = _control_key(control)
                    if key in clicked_controls:
                        continue
                    clicked_controls.add(key)
                    if len(clicked_controls - (baseline_controls or set())) > 4:
                        raise RuntimeError("A interface exibiu opções repetidas de exportação sem iniciar o download.")
                    LOGGER.info("Selecionando a opção da interface: %s", control.get("label", "download"))
                    _click(driver, control["element"])
                    ui_action = True
                    break
            if monitor is not None and monitor.file_url and not monitor.url_opened and now - (monitor.url_at or started) >= 3:
                _open_generated_download(driver, monitor)

        evidence = bool(candidates or partials or (monitor and (monitor.request_seen or monitor.download_started)))
        if driver is not None and monitor is not None and not evidence:
            if not ui_action and not retried_trigger and now - started >= 20:
                LOGGER.warning("Nenhuma ação de exportação observada; tentando uma vez o wrapper trigger do botão.")
                if monitor.main_handle in driver.window_handles:
                    driver.switch_to.window(monitor.main_handle)
                click_report_export(driver, wrapper=True)
                retried_trigger = True
            if now - started >= 60:
                raise TimeoutError("O clique não gerou menu de download, solicitação de exportação nem arquivo em 60s. Confira atms_diagnostic.json e atms_failure.png.")
        if now - last_progress >= 15:
            stage = monitor.stage if monitor is not None else "aguardando arquivo"
            LOGGER.info("Exportação ATMS: %s; tempo=%.0fs; arquivos=%s.", stage, now - started, sorted(changed))
            last_progress = now
        time.sleep(1)

    stage = monitor.stage if monitor is not None else "sem monitor da interface"
    raise TimeoutError(f"A exportação ATMS não concluiu em {timeout:g}s. Etapa: {stage}. "
                       f"Arquivos: {sorted(directory_snapshot(directory))}. "
                       "Consulte atms_diagnostic.json e atms_failure.png na pasta da execução.")


def normalize_header(value: Any) -> str:
    """Não modifica o arquivo: normaliza somente a comparação do cabeçalho."""
    if value is None:
        return ""
    text = str(value)
    for _ in range(2):
        text = html.unescape(text)
    text = unicodedata.normalize("NFC", text)
    text = text.replace("\ufeff", "").replace("\u200b", "")
    text = text.replace("*?*", "?")
    text = " ".join(text.split())
    return re.sub(r"\s+\?", "?", text)


def _worksheet_header(worksheet: Any, header_row: int) -> tuple[str, ...]:
    values = next(worksheet.iter_rows(
        min_row=header_row + 1, max_row=header_row + 1, values_only=True
    ), ())
    header = [normalize_header(value) for value in values]
    while header and not header[-1]:
        header.pop()
    return tuple(header)


def _schema_error(header: tuple[str, ...], sheet_name: str) -> ValueError:
    differences = []
    for index in range(max(len(header), len(EXPECTED_COLUMNS))):
        expected = EXPECTED_COLUMNS[index] if index < len(EXPECTED_COLUMNS) else "<sem coluna>"
        actual = header[index] if index < len(header) else "<ausente>"
        if expected != actual:
            differences.append(f"{index + 1}: esperado {expected!r}; encontrado {actual!r}")
    return ValueError(
        f"O schema ou a ordem do relatório ATMS mudou na aba {sheet_name!r}. "
        f"Esperadas {len(EXPECTED_COLUMNS)} colunas; encontradas {len(header)}. "
        + " | ".join(differences[:12])
        + (" | ..." if len(differences) > 12 else "")
    )


@dataclass(frozen=True)
class WorkbookValidation:
    sheet_name: str
    data_rows: int
    file_size: int


def validate_atms_workbook(
    path: Path, *, sheet: str = "auto", header_row: int = 0, allow_empty: bool = False
) -> WorkbookValidation:
    from openpyxl import load_workbook

    if header_row < 0:
        raise ValueError("--header-row deve ser zero ou maior (0 = primeira linha).")
    if not path.is_file():
        raise FileNotFoundError(f"XLSX não encontrado: {path}")
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"O arquivo deve ter extensão .xlsx: {path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError(f"O arquivo não é um XLSX válido: {path}") from exc
    try:
        if not workbook.worksheets:
            raise ValueError("O XLSX não possui planilhas.")
        # Algumas exportações descrevem dimensões imprecisas no XML.
        for worksheet in workbook.worksheets:
            worksheet.reset_dimensions()
        if sheet == "auto":
            matches = [ws for ws in workbook.worksheets
                       if _worksheet_header(ws, header_row) == EXPECTED_COLUMNS]
            if len(matches) > 1:
                raise ValueError("Várias abas têm o schema ATMS. Selecione uma com --sheet.")
            if not matches:
                ws = workbook["Report"] if "Report" in workbook.sheetnames else workbook.worksheets[0]
                raise _schema_error(_worksheet_header(ws, header_row), ws.title)
            worksheet = matches[0]
        elif re.fullmatch(r"\d+", sheet):
            index = int(sheet)
            if index >= len(workbook.worksheets):
                raise ValueError(f"Índice de aba inexistente: {index}.")
            worksheet = workbook.worksheets[index]
        else:
            if sheet not in workbook.sheetnames:
                raise ValueError(f"Aba {sheet!r} não encontrada. Abas: {workbook.sheetnames}.")
            worksheet = workbook[sheet]

        header = _worksheet_header(worksheet, header_row)
        if header != EXPECTED_COLUMNS:
            raise _schema_error(header, worksheet.title)
        data_rows = 0
        for row in worksheet.iter_rows(min_row=header_row + 2, values_only=True):
            if any(value not in (None, "") for value in row[len(EXPECTED_COLUMNS):]):
                raise ValueError("Há dados em colunas sem cabeçalho após a última coluna ATMS.")
            if any(value not in (None, "") for value in row[:len(EXPECTED_COLUMNS)]):
                data_rows += 1
        if not data_rows and not allow_empty:
            raise ValueError(
                "O relatório ATMS está vazio; o arquivo anterior não será substituído. "
                "Use --allow-empty somente se isso for esperado."
            )
        return WorkbookValidation(worksheet.title, data_rows, path.stat().st_size)
    finally:
        workbook.close()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def ensure_output_directory(directory: Path) -> None:
    if os.name == "nt" and str(directory).startswith("\\\\"):
        anchor = Path(directory.anchor)
        if not anchor.is_dir():
            raise OSError(f"Compartilhamento indisponível ou sem acesso: {anchor}")
    try:
        directory.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise OSError(f"Não foi possível acessar a pasta de destino {directory}: {exc}") from exc


def publish_atms_workbook(
    *, source: Path, output_dir: Path, sheet: str = "auto", header_row: int = 0,
    allow_empty: bool = False, lock_timeout: float = 30.0,
    replace_timeout: float = 30.0, delete_source: bool = False,
) -> tuple[Path, WorkbookValidation, bool]:
    if not source.is_file():
        raise FileNotFoundError(f"Arquivo de origem não encontrado: {source}")
    if source.suffix.lower() != ".xlsx":
        raise ValueError("A origem deve ser um arquivo .xlsx.")
    ensure_output_directory(output_dir)
    target = output_dir / FINAL_FILENAME
    if source.resolve() == target.resolve():
        validation = validate_atms_workbook(
            source, sheet=sheet, header_row=header_row, allow_empty=allow_empty
        )
        return target, validation, False

    with FileLock(Path(f"{target}.lock"), timeout=lock_timeout):
        replaced_existing = target.is_file()
        # Arquivo temporário no MESMO diretório final: evita shutil.move entre
        # discos e evita uma janela em que o destino não existiria.
        descriptor, name = tempfile.mkstemp(
            prefix=f".{target.stem}.", suffix=".tmp.xlsx", dir=output_dir
        )
        os.close(descriptor)
        temporary = Path(name)
        try:
            source_digest = _sha256(source)
            shutil.copyfile(source, temporary)
            with temporary.open("r+b") as stream:
                stream.flush()
                os.fsync(stream.fileno())
            if _sha256(temporary) != source_digest:
                raise OSError("A cópia do XLSX divergiu da origem; publicação cancelada.")
            # Valida precisamente os bytes que serão publicados na rede.
            validation = validate_atms_workbook(
                temporary, sheet=sheet, header_row=header_row, allow_empty=allow_empty
            )
            deadline = time.monotonic() + replace_timeout
            while True:
                try:
                    os.replace(temporary, target)
                    break
                except PermissionError as exc:
                    remaining = deadline - time.monotonic()
                    if remaining <= 0:
                        raise PermissionError(
                            f"Não foi possível substituir {target}. Feche o arquivo no Excel "
                            "e confira a permissão de alteração na pasta de rede. "
                            "O arquivo anterior foi preservado."
                        ) from exc
                    LOGGER.warning("Destino bloqueado; aguardando até %.0fs para substituir.", remaining)
                    time.sleep(min(2.0, remaining))
        finally:
            try:
                temporary.unlink(missing_ok=True)
            except OSError as exc:
                LOGGER.warning("Não foi possível remover o temporário %s: %s", temporary, exc)

    if delete_source:
        try:
            source.unlink(missing_ok=True)
        except OSError as exc:
            LOGGER.warning("Arquivo publicado; não foi possível remover o download %s: %s", source, exc)
    return target, validation, replaced_existing


def _browser_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--email", default=DEFAULT_EMAIL)
    parser.add_argument("--organization-url", default=DEFAULT_ORGANIZATION_URL)
    parser.add_argument("--report-url", default=DEFAULT_REPORT_URL)
    parser.add_argument("--user-data-dir", default=DEFAULT_EDGE_PROFILE)
    parser.add_argument("--step-timeout", type=int, default=120)
    parser.add_argument("--sso-timeout", type=int, default=300)
    parser.add_argument("--browser-start-timeout", type=int, default=900)


def _output_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR,
                        help="Pasta final; padrão: compartilhamento Report_Pipefy.")
    parser.add_argument("--sheet", default="auto",
                        help="auto seleciona a aba pelo schema; aceita nome ou índice a partir de 0.")
    parser.add_argument("--header-row", type=int, default=0,
                        help="Linha do cabeçalho a partir de 0 (padrão: primeira linha).")
    parser.add_argument("--allow-empty", action="store_true")
    parser.add_argument("--lock-timeout", type=float, default=30.0)
    parser.add_argument("--replace-timeout", type=float, default=30.0,
                        help="Espera por arquivo bloqueado no destino, em segundos.")


def _browser_kwargs(args: argparse.Namespace) -> dict[str, Any]:
    return dict(email=args.email, report_url=args.report_url,
                organization_url=args.organization_url,
                user_data_dir=args.user_data_dir, step_timeout=args.step_timeout,
                sso_timeout=args.sso_timeout, browser_start_timeout=args.browser_start_timeout)


def command_login(args: argparse.Namespace) -> None:
    _perform_visible_login(**_browser_kwargs(args))
    _mark_daily_login(args.user_data_dir)
    LOGGER.info("Perfil autenticado; relatório ATMS acessível.")


def command_check_headless(args: argparse.Namespace) -> None:
    driver, _ = _open_report_with_login_fallback(headless=True, **_browser_kwargs(args))
    driver.quit()
    LOGGER.info("Teste headless aprovado para o relatório ATMS.")


def _publish_from_args(args: argparse.Namespace, source: Path, *, delete_source: bool) -> None:
    target, validation, replaced = publish_atms_workbook(
        source=source, output_dir=resolve_path(args.output_dir), sheet=args.sheet,
        header_row=args.header_row, allow_empty=args.allow_empty,
        lock_timeout=args.lock_timeout, replace_timeout=args.replace_timeout,
        delete_source=delete_source,
    )
    LOGGER.info("%s publicado: linhas=%s, colunas=%s, aba=%s, bytes=%s, substituiu_anterior=%s.",
                target, validation.data_rows, len(EXPECTED_COLUMNS), validation.sheet_name,
                validation.file_size, replaced)


def command_run(args: argparse.Namespace) -> None:
    if args.weekdays_only and datetime.now().astimezone().weekday() >= 5:
        LOGGER.info("Execução ignorada: sábado/domingo.")
        return
    output_dir = resolve_path(args.output_dir)
    ensure_output_directory(output_dir)
    target = output_dir / FINAL_FILENAME
    staging_root = resolve_path(args.download_dir)
    staging_root.mkdir(parents=True, exist_ok=True)

    with FileLock(Path(f"{target}.run.lock"), timeout=args.lock_timeout):
        if args.ensure_daily_login:
            _ensure_daily_login(**_browser_kwargs(args))
        staging_dir = Path(tempfile.mkdtemp(prefix="run-", dir=staging_root))
        LOGGER.info("Exportando pipe=%s, report=%s. Download: %s", PIPE_ID, REPORT_ID, staging_dir)
        driver = None
        monitor = None
        before = directory_snapshot(staging_dir)
        try:
            try:
                driver, export_button = _open_report_with_login_fallback(
                    headless=args.headless, download_dir=staging_dir, capture_network=True,
                    **_browser_kwargs(args)
                )
                if not is_expected_report_url(driver.current_url):
                    raise RuntimeError("O navegador não está no relatório ATMS esperado.")
                monitor = ExportMonitor(driver)
                baseline_controls = {_control_key(c) for c in _download_controls(driver)}
                current_files = directory_snapshot(staging_dir)
                if any(is_expected_download_name(name) or name.endswith('.crdownload') for name in current_files):
                    LOGGER.info("[3/4] A página já iniciou o download nesta execução; aguardando conclusão.")
                else:
                    LOGGER.info("[3/4] Clicando no botão report-export e acompanhando a exportação da página.")
                    click_report_export(driver)
                downloaded = wait_for_download(
                    staging_dir, before, args.download_timeout, driver=driver,
                    monitor=monitor, baseline_controls=baseline_controls,
                )
                LOGGER.info("Download concluído: %s", downloaded)
            except Exception:
                if driver is not None:
                    save_export_diagnostics(driver, staging_dir, monitor)
                raise
            finally:
                if driver is not None:
                    try:
                        driver.quit()
                    except Exception as exc:
                        LOGGER.warning("Falha ao encerrar o Edge: %s", exc)
            LOGGER.info("[4/4] Validando o XLSX e substituindo o arquivo de destino.")
            _publish_from_args(args, downloaded, delete_source=True)
        except Exception:
            LOGGER.error("Execução falhou. Downloads disponíveis para diagnóstico em: %s", staging_dir)
            raise
        else:
            try:
                staging_dir.rmdir()
            except OSError:
                pass


def command_replace(args: argparse.Namespace) -> None:
    _publish_from_args(args, resolve_path(args.input), delete_source=False)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Exporta Status Projeto ATMS pelo Edge e substitui status_projeto_atms.xlsx."
    )
    parser.add_argument("--version", action="version", version=SCRIPT_VERSION)
    parser.add_argument("--verbose", action="store_true")
    parser.add_argument("--silent", action="store_true", help="Mantém apenas o log em arquivo.")
    parser.add_argument("--log-file", default=DEFAULT_LOG_FILE)
    subparsers = parser.add_subparsers(dest="command", required=True)
    login = subparsers.add_parser("login", help="Autentica SSO com janela visível.")
    _browser_arguments(login)
    login.set_defaults(func=command_login)
    check = subparsers.add_parser("check-headless", help="Confere acesso sem exportar.")
    _browser_arguments(check)
    check.set_defaults(func=command_check_headless)
    run = subparsers.add_parser("run", aliases=["flow"], help="Baixa, valida e substitui o XLSX.")
    _browser_arguments(run)
    _output_arguments(run)
    run.add_argument("--headless", action="store_true")
    run.add_argument("--weekdays-only", action="store_true")
    run.add_argument("--ensure-daily-login", action="store_true")
    run.add_argument("--download-timeout", type=int, default=600)
    run.add_argument("--download-dir", default=DEFAULT_DOWNLOAD_DIR,
                     help="Raiz temporária local; cada execução usa uma subpasta isolada.")
    run.set_defaults(func=command_run)
    replace = subparsers.add_parser("replace", help="Valida e publica um XLSX já baixado.")
    replace.add_argument("--input", required=True)
    _output_arguments(replace)
    replace.set_defaults(func=command_replace)
    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()
    for field in ("step_timeout", "sso_timeout", "browser_start_timeout", "download_timeout"):
        if hasattr(args, field) and getattr(args, field) <= 0:
            parser.error(f"--{field.replace('_', '-')} deve ser maior que zero.")
    for field in ("header_row", "lock_timeout", "replace_timeout"):
        if hasattr(args, field) and getattr(args, field) < 0:
            parser.error(f"--{field.replace('_', '-')} não pode ser negativo.")
    try:
        configure_logging(args.verbose, args.log_file, args.silent)
        LOGGER.info("Status ATMS v%s: organização -> relatório -> exportação -> arquivo.", SCRIPT_VERSION)
        args.func(args)
    except KeyboardInterrupt:
        LOGGER.error("Execução cancelada.")
        raise SystemExit(130)
    except Exception as exc:
        if args.verbose:
            LOGGER.exception("Falha: %s", exc)
        else:
            LOGGER.error("Falha: %s", exc)
        raise SystemExit(1)


if __name__ == "__main__":
    main()
