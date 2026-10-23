from __future__ import annotations

import contextlib
import csv
import datetime as dt
import hashlib
import os
import shutil
import tempfile
import uuid
import zipfile
from pathlib import Path
from typing import Any, Callable, Iterable

from .io import Lock, atomic_json, digest, read_json


class EventConflict(ValueError):
    def __init__(self, conflicts: list[dict[str, Any]], total: int):
        self.conflicts = conflicts
        super().__init__(
            f"Eventos: {total} chaves com conteúdo divergente. Histórico preservado; revise os conflitos antes de publicar."
        )


def xlsx_content(path: Path) -> None:
    from openpyxl import load_workbook

    if path.suffix.lower() != ".xlsx" or not zipfile.is_zipfile(path):
        raise ValueError("Arquivo não é um XLSX íntegro.")
    with zipfile.ZipFile(path) as archive:
        if archive.testzip():
            raise ValueError("CRC inválido no XLSX.")
    book = load_workbook(path, read_only=True)
    try:
        for sheet in book:
            count = 0
            for row in sheet.iter_rows(values_only=True):
                if any(v is not None and str(v).strip() for v in row):
                    count += 1
                if count >= 2:
                    return
        raise ValueError("Arquivo vazio ou sem linhas de dados.")
    finally:
        book.close()


def synced_copy(source: Path, target: Path) -> None:
    with source.open("rb") as incoming, target.open("wb") as outgoing:
        shutil.copyfileobj(incoming, outgoing, 1024 * 1024)
        outgoing.flush()
        os.fsync(outgoing.fileno())
    if digest(source) != digest(target):
        raise OSError("Falha de integridade ao copiar o arquivo.")


class Publisher:
    def __init__(self, backups: Path, retention: dict[str, Any]):
        self.backups = backups
        self.retention = retention

    @staticmethod
    def lock_path(target: Path) -> Path:
        key = hashlib.sha256(str(target.resolve()).encode()).hexdigest()[:20]
        return target.parent / f".orquestrador-{key}.lock"

    def backup(self, target: Path, run: str) -> dict[str, Any] | None:
        if not target.is_file():
            return None
        self.backups.mkdir(parents=True, exist_ok=True)
        identifier = str(uuid.uuid4())
        content = self.backups / f"{identifier}.data"
        synced_copy(target, content)
        record = {
            "owner": "orquestrador-report-v1",
            "id": identifier,
            "target": str(target.resolve()),
            "at": dt.datetime.now(dt.timezone.utc).isoformat(),
            "run": run,
            "size": content.stat().st_size,
            "sha256": digest(content),
        }
        try:
            atomic_json(self.backups / f"{identifier}.json", record)
        except Exception:
            content.unlink(missing_ok=True)
            raise
        return record

    def versions(self, target: Path | None = None) -> list[dict[str, Any]]:
        result = []
        for path in self.backups.glob("*.json"):
            if path.is_symlink():
                continue
            try:
                record = read_json(path)
                if record.get("owner") != "orquestrador-report-v1" or path.stem != record["id"]:
                    continue
                if target is None or record["target"] == str(target.resolve()):
                    result.append(record)
            except (OSError, ValueError, KeyError):
                continue
        return sorted(result, key=lambda r: r["at"], reverse=True)

    def validate_version(self, record: dict[str, Any]) -> Path:
        identifier = record["id"]
        if str(uuid.UUID(identifier)) != identifier:
            raise ValueError("Identificador de versão inválido.")
        content = self.backups / f"{identifier}.data"
        if (
            content.is_symlink()
            or content.stat().st_size != record["size"]
            or digest(content) != record["sha256"]
        ):
            raise ValueError("Backup inválido: tamanho/hash divergentes.")
        return content

    def _replace(self, source: Path, target: Path, run: str) -> dict[str, Any]:
        fd, name = tempfile.mkstemp(
            prefix=".orquestrador-publish-", suffix=target.suffix, dir=target.parent
        )
        os.close(fd)
        temporary = Path(name)
        try:
            synced_copy(source, temporary)
            backup = self.backup(target, run)
            os.replace(temporary, target)
            result = {
                "file": str(target),
                "size": target.stat().st_size,
                "sha256": digest(target),
                "backup": backup,
                "publication": "local; sincronização externa não verificada",
            }
        finally:
            temporary.unlink(missing_ok=True)
        # Retention failures are maintenance issues after a committed publication, never a false extraction failure.
        try:
            self.prune(target)
        except (OSError, ValueError) as exc:
            result["retention_warning"] = str(exc)
        return result

    def publish(
        self, source: Path, target: Path, run: str, validator: Callable[[Path], Any] = xlsx_content
    ) -> dict[str, Any]:
        if source.suffix != target.suffix:
            raise ValueError("Extensão da origem difere do contrato de destino.")
        if any(p.is_file() for p in source.parent.glob(source.name + ".crdownload")):
            raise ValueError("Download parcial não pode ser publicado.")
        validator(source)
        target.parent.mkdir(parents=True, exist_ok=True)
        with Lock(self.lock_path(target)):
            return self._replace(source, target, run)

    def restore(self, identifier: str, run: str) -> dict[str, Any]:
        record = next((r for r in self.versions() if r["id"] == identifier), None)
        if not record:
            raise ValueError("Versão de relatório não encontrada.")
        content = self.validate_version(record)
        target = Path(record["target"])
        if not target.parent.is_dir():
            raise ValueError("Pasta original indisponível.")
        with Lock(self.lock_path(target)):
            result = self._replace(content, target, run)
        result["operation"] = "restore"
        return result

    def prune(self, target: Path) -> None:
        records = self.versions(target)
        now = dt.datetime.now(dt.timezone.utc)
        total = 0
        for index, record in enumerate(records):
            total += record["size"]
            age = (now - dt.datetime.fromisoformat(record["at"])).total_seconds() / 86400
            remove = index > 0 and (
                index >= int(self.retention["count"])
                or age > int(self.retention["days"])
                or total > int(self.retention["bytes"])
            )
            if remove:
                self.validate_version(record)
                (self.backups / f"{record['id']}.json").unlink()
                (self.backups / f"{record['id']}.data").unlink()

    def merge_rows(
        self,
        rows: Iterable[dict[str, Any]],
        columns: list[str],
        target: Path,
        run: str,
        keys: list[str],
        separator: str,
        encoding: str,
    ) -> dict[str, Any]:
        if not set(keys) <= set(columns):
            raise ValueError("Chave ausente nas colunas de Eventos.")
        target.parent.mkdir(parents=True, exist_ok=True)
        with Lock(self.lock_path(target)):
            fd, name = tempfile.mkstemp(
                prefix=".orquestrador-events-", suffix=".csv", dir=target.parent
            )
            os.close(fd)
            temporary = Path(name)
            seen: dict[tuple[str, ...], tuple[str, ...]] = {}
            conflicts: list[dict[str, Any]] = []
            conflict_count = 0
            counts = {"incoming": 0, "duplicates": 0, "appended": 0}

            def key_for(row: dict[str, Any], line: int) -> tuple[str, ...]:
                values = tuple(
                    str(row.get(k) if row.get(k) is not None else "").strip() for k in keys
                )
                if any(not value or value.casefold() in {"nan", "none", "nat"} for value in values):
                    raise ValueError(
                        f"Chave de Eventos ausente/inválida na linha {line}; histórico preservado."
                    )
                return values

            try:
                with temporary.open("w", encoding=encoding, newline="") as output:
                    writer = csv.DictWriter(
                        output, fieldnames=columns, delimiter=separator, lineterminator="\n"
                    )
                    writer.writeheader()
                    if target.exists():
                        with target.open(encoding=encoding, newline="") as old:
                            reader = csv.DictReader(old, delimiter=separator)
                            if reader.fieldnames != columns:
                                raise ValueError(
                                    "Colunas ou ordem do histórico de Eventos divergentes."
                                )
                            for line, row in enumerate(reader, 2):
                                key = key_for(row, line)
                                old_content = tuple(
                                    str(row.get(c) if row.get(c) is not None else "")
                                    for c in columns
                                )
                                if key in seen and seen[key] != old_content:
                                    raise EventConflict(
                                        [
                                            {
                                                "key": dict(zip(keys, key)),
                                                "existing": dict(zip(columns, seen[key])),
                                                "incoming": row,
                                                "line": line,
                                            }
                                        ],
                                        1,
                                    )
                                seen[key] = old_content
                                writer.writerow(row)
                    for line, row in enumerate(rows, 2):
                        counts["incoming"] += 1
                        key = key_for(row, line)
                        content = tuple(
                            str(row.get(c) if row.get(c) is not None else "") for c in columns
                        )
                        if key in seen:
                            if seen[key] != content:
                                conflict_count += 1
                                if len(conflicts) < 100:
                                    conflicts.append(
                                        {
                                            "key": dict(zip(keys, key)),
                                            "before": dict(zip(columns, seen[key])),
                                            "incoming": dict(zip(columns, content)),
                                        }
                                    )
                            counts["duplicates"] += 1
                            continue
                        seen[key] = content
                        writer.writerow(row)
                        counts["appended"] += 1
                    output.flush()
                    os.fsync(output.fileno())
                if counts["incoming"] == 0:
                    raise ValueError("Exportação de Eventos sem dados.")
                if conflict_count:
                    raise EventConflict(conflicts, conflict_count)
                if counts["appended"] == 0 and target.exists():
                    return {
                        **counts,
                        "file": str(target),
                        "sha256": digest(target),
                        "unchanged": True,
                    }
                return {**counts, **self._replace(temporary, target, run)}
            finally:
                temporary.unlink(missing_ok=True)


def cleanup_partials(folder: Path, older_days: int) -> list[str]:
    """Explicit maintenance only, confined to managed run-* directories."""
    removed = []
    threshold = dt.datetime.now().timestamp() - older_days * 86400
    for path in folder.glob("run-*"):
        if path.is_dir() and not path.is_symlink() and path.stat().st_mtime < threshold:
            with contextlib.suppress(OSError):
                shutil.rmtree(path)
                removed.append(path.name)
    return removed
