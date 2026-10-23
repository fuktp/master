"""Sanitized JSON/XLSX output, without loading pandas into the UI."""

from __future__ import annotations

import json
import os
import tempfile
from pathlib import Path
from typing import Any

from .io import atomic_json, sanitize


def export_data(target: Path, value: Any) -> dict[str, str]:
    value = sanitize(value, export=True)
    if target.suffix.casefold() != ".xlsx":
        atomic_json(target, value)
    else:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell

        workbook = Workbook(write_only=True)
        overflow = None
        groups = value if isinstance(value, dict) else {"history": value}
        for index, (name, content) in enumerate(groups.items()):
            sheet = workbook.create_sheet(f"{index + 1}-{name}"[:31].replace("/", "-"))
            rows = content if isinstance(content, list) else [content]
            records = [r if isinstance(r, dict) else {"value": r} for r in rows]
            if len(records) >= 1_048_575:
                raise ValueError(
                    "O conteúdo excede o limite de linhas do XLSX; exporte em JSON, sem truncamento."
                )
            keys = list(dict.fromkeys(k for record in records for k in record)) or ["value"]
            sheet.append(keys)
            for row_number, record in enumerate(records, 2):
                cells = []
                for key in keys:
                    cell = record.get(key)
                    if isinstance(cell, (dict, list)):
                        cell = json.dumps(cell, ensure_ascii=False)
                    if isinstance(cell, str):
                        cell = "".join(
                            c if c in "\t\r\n" or ord(c) >= 32 else f"\\u{ord(c):04x}" for c in cell
                        )
                        if len(cell.encode("utf-16-le")) > 64000:
                            if overflow is None:
                                overflow = workbook.create_sheet("Conteúdo longo")
                                overflow.append(
                                    ["Planilha", "Linha", "Campo", "Parte", "Total", "Conteúdo"]
                                )
                            chunks = [cell[i : i + 16000] for i in range(0, len(cell), 16000)]
                            for part, chunk in enumerate(chunks, 1):
                                text_cell = WriteOnlyCell(overflow, value=chunk)
                                text_cell.data_type = "s"
                                overflow.append(
                                    [sheet.title, row_number, key, part, len(chunks), text_cell]
                                )
                            cell = f"[Conteúdo completo: planilha Conteúdo longo; linha {row_number}; campo {key}; {len(chunks)} partes]"
                        text_cell = WriteOnlyCell(sheet, value=cell)
                        text_cell.data_type = "s"
                        cell = text_cell
                    cells.append(cell)
                sheet.append(cells)
        target.parent.mkdir(parents=True, exist_ok=True)
        fd, name = tempfile.mkstemp(suffix=".xlsx", dir=target.parent)
        os.close(fd)
        try:
            workbook.save(name)
            os.replace(name, target)
        finally:
            Path(name).unlink(missing_ok=True)
            workbook.close()
    return {"file": str(target), "status": "exported"}
