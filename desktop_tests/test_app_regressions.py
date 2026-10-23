from __future__ import annotations

import json
import shutil
import sys
import time
from pathlib import Path
from types import ModuleType, SimpleNamespace
from typing import Any

import pytest
from PySide6.QtCore import QItemSelectionModel
from PySide6.QtGui import QColor, QPalette
from PySide6.QtWidgets import QApplication, QCheckBox, QFileDialog, QMessageBox

from desktop_tests.ui_helpers import settle
from orquestrador_app.config import Config, Routine, defaults, load, portable, save
from orquestrador_app.gui import MainWindow, RoutineDialog
from orquestrador_app.io import Lock, read_json
from orquestrador_app.paths import Paths
from orquestrador_app.resources import asset_path, source_root
from orquestrador_app.store import Store
from orquestrador_app.ui import icon, theme


@pytest.fixture
def window(tmp_path: Path) -> Any:
    app = QApplication.instance() or QApplication([])
    dark = QPalette()
    dark.setColor(QPalette.ColorRole.Text, QColor("white"))
    dark.setColor(QPalette.ColorRole.WindowText, QColor("white"))
    dark.setColor(QPalette.ColorRole.Base, QColor("white"))
    app.setPalette(dark)
    paths = Paths.user(tmp_path / "app")
    paths.ensure()
    config = defaults(paths, source_root())
    save(paths, config)
    result = MainWindow(paths)
    settle(result)
    result.resize(1280, 850)
    result.show()
    app.processEvents()
    try:
        yield result
    finally:
        if result.task and result.task.isRunning():
            result.task.wait(5000)
        result.exit_requested = True
        result.close()
        app.processEvents()


def drain(window: MainWindow) -> None:
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        QApplication.processEvents()
        if window.task and not window.task.isRunning():
            QApplication.processEvents()
            return
        time.sleep(0.005)
    pytest.fail("A operação da interface não concluiu em cinco segundos")


def test_all_pages_keep_activity_visible_and_render_dark_cell_text(window: MainWindow) -> None:
    assert window.width() <= 1280, "Os textos não devem forçar a janela para fora de uma tela menor"
    for index in range(window.tabs.count()):
        window.tabs.setCurrentIndex(index)
        QApplication.processEvents()
        assert window.log.isVisible() and window.status_label.isVisible()
    window.tabs.setCurrentIndex(0)
    QApplication.processEvents()
    table = window.panel
    assert table.item(0, 0).foreground().color().name() == theme()["foreground"].lower()
    rect = table.visualItemRect(table.item(0, 0)).adjusted(5, 5, -5, -5)
    image = table.viewport().grab(rect).toImage()
    ink = sum(
        max(
            image.pixelColor(x, y).red(),
            image.pixelColor(x, y).green(),
            image.pixelColor(x, y).blue(),
        )
        < 110
        for y in range(image.height())
        for x in range(image.width())
    )
    assert ink > 10, "A célula deve desenhar texto escuro sem precisar ser selecionada"
    for group in (QPalette.ColorGroup.Active, QPalette.ColorGroup.Inactive):
        assert table.palette().color(group, QPalette.ColorRole.Text).name() == "#333333"
    dialog = RoutineDialog(window.config, window.config.routines[0], window)
    try:
        dialog.show()
        QApplication.processEvents()
        assert all(box.text() for box in dialog.findChildren(QCheckBox))
        assert dialog.script.palette().color(QPalette.ColorRole.Text).name() == "#333333"
        assert dialog.script.palette().color(QPalette.ColorRole.Base).name() == "#ffffff"
        assert all(button.text() and button.toolTip() for button in dialog.days_picker.buttons)
    finally:
        dialog.reject()


def test_icon_is_shared_by_window_tray_and_notification(
    window: MainWindow, monkeypatch: pytest.MonkeyPatch
) -> None:
    from orquestrador_app.notifications import show

    def pixels(value: Any) -> bytes:
        return bytes(value.pixmap(32, 32).toImage().constBits())

    assert pixels(window.windowIcon()) == pixels(window.tray.icon()) == pixels(icon())
    calls = []
    show(SimpleNamespace(showMessage=lambda *args: calls.append(args)), "Teste local")
    assert calls[0][:2] == ("CAIO", "Teste local")
    assert pixels(calls[0][2]) == pixels(icon())


def test_selection_groups_and_cancel_only_selected_execution(window: MainWindow) -> None:
    assert window.batch_group.isVisible()
    assert not window.run_button.isEnabled()
    window.panel.selectRow(0)
    assert not window.batch_group.isVisible() and window.run_button.isEnabled()
    selection = window.panel.selectionModel()
    selection.select(
        window.panel.model().index(1, 0),
        QItemSelectionModel.SelectionFlag.Select | QItemSelectionModel.SelectionFlag.Rows,
    )
    assert "lote" in window.run_button.text().lower()
    window.panel.clearSelection()
    window.panel.selectRow(0)
    first, second = [window.panel.item(i, 0).data(256) for i in range(2)]
    one = window.store.begin("b", first, "extraction", None)
    two = window.store.begin("b", second, "extraction", None)
    window.control("cancel")
    drain(window)
    assert window.store.control(one) == "cancel"
    assert window.store.control(two) != "cancel"
    window.store.finish(one, "cancelled", 130)
    window.store.finish(two, "cancelled", 130)


def test_second_launch_requests_foreground_and_existing_window_responds(
    window: MainWindow, monkeypatch: pytest.MonkeyPatch
) -> None:
    from orquestrador_app.__main__ import main

    monkeypatch.setattr("orquestrador_app.native_updates.is_native", lambda _: True)
    window.hide()
    with Lock(window.paths.root / "state/interface.lock"):
        assert main(["--data-root", str(window.paths.root)]) == 0
    assert (window.paths.root / "state/open-interface.json").is_file()
    window.refresh()
    settle(window)
    assert window.isVisible()
    assert not (window.paths.root / "state/open-interface.json").exists()


def test_manual_force_on_paused_routine_uses_changed_script_and_records_success(
    tmp_path: Path,
) -> None:
    from orquestrador_app.engine import Engine

    paths = Paths.user(tmp_path / "app")
    paths.ensure()
    original = tmp_path / "original.py"
    replacement = tmp_path / "revised.py"
    original.write_text('print("original")')
    replacement.write_text('print("revised")')
    config = defaults(paths, source_root())
    job = Routine("additional", "Additional", str(original), requires_browser=False, command="")
    job.schedule.paused = True
    job.schedule.days = list(range(7))
    config.routines = [job]
    config.scheduling_paused = True
    job.script = str(replacement)
    save(paths, config)
    result = Engine(paths).run([job.id], force=True)
    assert result[0]["status"] == "success"
    assert Store(paths.database).state(job.id)["last_success"]
    assert Engine(paths).run([job.id], scheduled=True)[0]["status"] == "not_eligible"
    assert original.read_text() == 'print("original")'


def test_full_driver_repair_copies_local_compatible_version_and_keeps_backup(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from orquestrador_app.driver import guard

    paths = Paths.user(tmp_path / "app")
    paths.ensure()
    local = tmp_path / "tools/msedgedriver.exe"
    local.parent.mkdir()
    local.write_text("140.0.10.7")
    target = paths.root / "cache/driver/msedgedriver.exe"
    target.parent.mkdir(parents=True)
    target.write_text("139.0.10.1")
    config = Config(
        driver=str(target),
        logs=str(paths.root / "logs"),
        driver_options={"search_paths": [str(local)]},
    )
    instance = guard(paths, config)
    browser = tmp_path / "msedge.exe"
    browser.write_text("fixture")
    monkeypatch.setattr(instance, "detect_edge_binary", lambda: browser)
    monkeypatch.setattr(instance, "edge_version", lambda *_: (140, 0, 10, 99))
    monkeypatch.setattr(
        instance, "driver_version", lambda path: tuple(map(int, path.read_text().split(".")))
    )
    monkeypatch.setattr(instance, "authenticode_status", lambda _: "Valid")
    details = instance.ensure(repair=True, test_session=False)
    assert details["compatible"] and target.read_text() == local.read_text()
    assert next(target.parent.glob("*.backup-*")).read_text() == "139.0.10.1"
    assert local.read_text() == "140.0.10.7"


def test_driver_reads_literal_path_from_registered_script_without_importing_it(
    tmp_path: Path,
) -> None:
    from orquestrador_app.driver import candidates

    external = tmp_path / "outside/msedgedriver.exe"
    external.parent.mkdir()
    external.write_bytes(b"fixture")
    script = tmp_path / "extract.py"
    script.write_text("raise RuntimeError('never import')\nDRIVER = " + repr(str(external)))
    config = Config(routines=[Routine("x", "x", str(script))])
    assert external.resolve() in candidates(Paths.user(tmp_path / "app"), config)


@pytest.mark.parametrize(
    "outer,inner",
    [
        (0x80070002, None),
        (-2147024894, None),
        (-2147352567, -2147024894),
        (0x80020009, 0x80070002),
    ],
)
def test_task_com_query_distinguishes_missing_from_denied(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, outer: int, inner: int | None
) -> None:
    import orquestrador_app.tasks as tasks

    class ComError(Exception):
        def __init__(self, code: int, detail: int | None = None):
            self.hresult = code
            self.excepinfo = (0, None, None, None, 0, detail) if detail is not None else None
            super().__init__(code, "Exception occurred.", self.excepinfo, None)

    value: list[Any] = [ComError(outer, inner)]

    def get_task(_: str) -> Any:
        if isinstance(value[0], Exception):
            raise value[0]
        return SimpleNamespace(Xml=value[0])

    service = SimpleNamespace(
        Connect=lambda: None, GetFolder=lambda _: SimpleNamespace(GetTask=get_task)
    )
    client = SimpleNamespace(Dispatch=lambda _: service)
    package = ModuleType("win32com")
    package.client = client
    monkeypatch.setitem(sys.modules, "win32com", package)
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    lifecycle = [0, 0]
    monkeypatch.setitem(
        sys.modules,
        "pythoncom",
        SimpleNamespace(
            CoInitialize=lambda: lifecycle.__setitem__(0, lifecycle[0] + 1),
            CoUninitialize=lambda: lifecycle.__setitem__(1, lifecycle[1] + 1),
        ),
    )
    monkeypatch.setitem(sys.modules, "pywintypes", SimpleNamespace(com_error=ComError))
    monkeypatch.setattr(tasks, "os", SimpleNamespace(name="nt"))
    backend = tasks.WindowsTasks(Paths.user(tmp_path))
    assert backend.export() is None
    from orquestrador_app.native_install import NativeInstaller

    installer = NativeInstaller(
        backend.paths, tasks=backend, links=SimpleNamespace(snapshot=lambda _: None)
    )
    assert installer.snapshot(tmp_path / "backup")["had_task"] is False
    for error in (
        ComError(0x80070005),
        ComError(-2147352567, -2147024891),
        ComError(0x8004130F),
        ComError(0x80020009, 0x8004130F),
        ComError(-2147352567),
    ):
        value[0] = error
        with pytest.raises(RuntimeError, match="consultar"):
            backend.export()
    value[0] = "<Task />"
    assert backend.export() == "<Task />"

    def missing_folder(_: str) -> Any:
        raise ComError(0x80070002)

    service.GetFolder = missing_folder
    with pytest.raises(ComError):
        backend.export()
    assert lifecycle[0] == lifecycle[1] == 9


@pytest.mark.parametrize("frequency", ["daily", "weekly", "every_run", "weekdays"])
def test_editor_calendar_custom_detection_and_frequency_presets(
    frequency: str, tmp_path: Path
) -> None:
    from orquestrador_app.config import Schedule, preset_days

    app = QApplication.instance() or QApplication([])
    script = tmp_path / "fixture.py"
    script.write_text("# Local fixture; not executed by this dialog test.\n")
    job = Routine(
        "local", "Local", str(script), schedule=Schedule(frequency=frequency, days=[1, 3])
    )
    dialog = RoutineDialog(Config(routines=[job]), job)
    try:
        assert dialog.day_pattern.currentText() == "Personalizado"
        assert dialog.days_picker.values() == [1, 3]
        assert dialog.frequency.currentData() == frequency
        dialog.day_pattern.activated.emit(0)
        assert dialog.days_picker.values() == preset_days(frequency)
        assert dialog.day_pattern.currentIndex() == 0
        dialog.days_picker.buttons[2].toggle()
        assert dialog.day_pattern.currentIndex() == 1
        assert dialog.frequency.currentData() == frequency
        dialog.times.clock.setTime(dialog.times.clock.time().addSecs(3600))
        dialog.times.add()
        assert dialog.times.values() == ["08:30", "09:30"]
        dialog.accept_values()
        assert dialog.updated.schedule.days == dialog.days_picker.values()
        assert dialog.updated.schedule.frequency == frequency
        app.processEvents()
    finally:
        dialog.done(0)


def test_windows_shortcut_and_window_use_same_app_id_and_hashed_icon(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from orquestrador_app.native_install import WindowsLinks
    from orquestrador_app.windows_identity import app_id, launch_icon, set_window_identity

    paths = Paths.user(tmp_path)
    paths.ensure()
    (paths.root / "launch").mkdir()
    (paths.root / "launch/Iniciar.vbs").write_text("' fixture")
    branded = launch_icon(paths)
    shutil.copy2(asset_path("caio.ico"), branded)
    shortcut = SimpleNamespace(TargetPath="", Save=lambda: None)
    stores: dict[str, dict[str, Any]] = {"shortcut": {}, "window": {}}

    def property_store(name: str) -> Any:
        return SimpleNamespace(
            SetValue=lambda key, value: stores[name].__setitem__(key, value),
            Commit=lambda: None,
        )

    client = SimpleNamespace(Dispatch=lambda _: SimpleNamespace(CreateShortcut=lambda _: shortcut))
    package = ModuleType("win32com")
    package.client = client
    propsys = SimpleNamespace(
        SHGetPropertyStoreFromParsingName=lambda *args: property_store("shortcut"),
        SHGetPropertyStoreForWindow=lambda *args: property_store("window"),
        PROPVARIANTType=lambda value: value,
    )
    pscon = SimpleNamespace(
        PKEY_AppUserModel_ID="id",
        PKEY_AppUserModel_RelaunchCommand="command",
        PKEY_AppUserModel_RelaunchIconResource="icon",
        PKEY_AppUserModel_RelaunchDisplayNameResource="name",
    )
    monkeypatch.setitem(sys.modules, "win32com", package)
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    monkeypatch.setitem(
        sys.modules, "win32com.propsys", SimpleNamespace(propsys=propsys, pscon=pscon)
    )
    monkeypatch.setitem(
        sys.modules, "win32com.shell", SimpleNamespace(shellcon=SimpleNamespace(GPS_READWRITE=2))
    )
    monkeypatch.setitem(
        sys.modules,
        "pythoncom",
        SimpleNamespace(CoInitialize=lambda: None, CoUninitialize=lambda: None),
    )
    links = WindowsLinks(paths)
    monkeypatch.setattr(links, "locations", lambda: [tmp_path / "CAIO.lnk"])
    links.install()
    monkeypatch.setattr("orquestrador_app.windows_identity.sys.platform", "win32")
    monkeypatch.setenv("WINDIR", "C:/Windows")
    set_window_identity(SimpleNamespace(winId=lambda: 123), paths)
    assert stores["shortcut"]["id"] == stores["window"]["id"] == app_id(paths)
    assert shortcut.IconLocation == stores["window"]["icon"] == str(branded) + ",0"
    assert "wscript.exe" in stores["window"]["command"]
    assert "Iniciar.vbs" in stores["window"]["command"]


def test_task_results_populate_comparison_in_app(window: MainWindow) -> None:
    from orquestrador_app.tasks import xml

    expected = xml(Path("Iniciar.vbs"), window.paths, "user")
    window.tabs.setCurrentIndex(4)
    window.task_done(
        {"installed": True, "matches": True, "expected": expected, "actual": expected}, False
    )
    assert window.tasks_view.rowCount() == 6
    assert all(window.tasks_view.item(i, 3).text() == "OK" for i in range(6))
    assert "null" not in window.log.toPlainText()


def test_task_table_uses_effective_values_for_equivalent_xml(
    window: MainWindow, normalized_task_xml: Any
) -> None:
    from orquestrador_app.tasks import compare_tasks, xml

    expected = xml(Path("Iniciar.vbs"), window.paths, "fixture")
    report = compare_tasks(normalized_task_xml(expected), expected)
    assert report["matches"]
    window.task_done(report, False)
    assert window.tasks_view.rowCount() == len(report["checks"])
    assert all(
        window.tasks_view.item(i, 3).text() == "OK" for i in range(window.tasks_view.rowCount())
    )
    report = compare_tasks(normalized_task_xml(expected).replace("PT60S", "PT2M"), expected)
    window.task_done(report, False)
    assert window.status_label.property("severity") == "warning"


def test_preference_export_import_uses_actual_ui_callbacks(
    window: MainWindow, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    file = tmp_path / "prefs.json"
    expected = portable(load(window.paths))
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: (str(file), "JSON"))
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *args, **kwargs: (str(file), "JSON"))
    monkeypatch.setattr("orquestrador_app.gui.show_detail", lambda *args: None)
    monkeypatch.setattr(QMessageBox, "question", lambda *args: QMessageBox.StandardButton.Yes)
    window.export_settings()
    config = load(window.paths)
    config.queue_delay_seconds = 99
    save(window.paths, config)
    window.refresh()
    settle(window)
    window.import_settings()
    assert portable(load(window.paths)) == expected == read_json(file)


@pytest.mark.parametrize("suffix", ["xlsx", "json"])
def test_history_export_actual_callback(
    window: MainWindow, tmp_path: Path, suffix: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    output = tmp_path / ("history." + suffix)
    run = window.store.begin("synthetic", window.config.routines[0].id, "extraction", None)
    window.store.finish(run, "failed", 1, evidence="fixture local")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: (str(output), ""))
    window.export_history()
    drain(window)
    assert output.is_file() and output.stat().st_size > 0
    if suffix == "xlsx":
        from openpyxl import load_workbook

        book = load_workbook(output, read_only=True)
        assert book.sheetnames
        book.close()
    else:
        values = json.loads(output.read_text())
        assert isinstance(values, list) and len(values) == 1


def test_full_theme_is_durable_and_contrast_is_readable() -> None:
    values = read_json(asset_path("dashboard-theme.json"))
    assert values["name"] == "Dashboard de Produtividade - KPMG Claro"
    assert len(values["dataColors"]) == 12
    assert set(values["visualStyles"]) == {
        "*",
        "page",
        "card",
        "cardVisual",
        "slicer",
        "tableEx",
        "pivotTable",
    }
    assert values["textClasses"]["largeTitle"]["fontFace"] == "Segoe UI Semibold"
