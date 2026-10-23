from __future__ import annotations

import datetime as dt
import os
import sys
from pathlib import Path
from types import ModuleType, SimpleNamespace
from typing import Any

import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QMessageBox

from desktop_tests.ui_helpers import settle
from orquestrador_app.config import (
    Config,
    ConfigurationConflict,
    Dependency,
    Routine,
    Schedule,
    defaults,
    load,
    save,
    validate,
)
from orquestrador_app.engine import Engine
from orquestrador_app.exporting import export_data
from orquestrador_app.files import EventConflict, Publisher
from orquestrador_app.gui import MainWindow, RoutineDialog
from orquestrador_app.io import digest, read_json
from orquestrador_app.monitor_status import monitor_state, poll_seconds, write_heartbeat
from orquestrador_app.notifications import claim, enqueue
from orquestrador_app.observation import HistoryModel, LogTail
from orquestrador_app.operations import (
    clean_previewed,
    export_routine,
    export_routine_preview,
    quarantine,
    storage_preview,
)
from orquestrador_app.paths import Paths
from orquestrador_app.resources import source_root
from orquestrador_app.schedule import eligible, next_execution, period_for
from orquestrador_app.script_versions import (
    capture,
    check_dependencies,
    remove,
    restore,
    restore_registration,
)
from orquestrador_app.store import Store
from orquestrador_app.tasks import NS, WindowsTasks

UTC = dt.timezone.utc
NOW = dt.datetime(2026, 10, 19, 9, tzinfo=UTC)  # Monday


def local_job(tmp_path: Path) -> tuple[Paths, Config, Routine]:
    paths = Paths.user(tmp_path / "data")
    paths.ensure()
    source = tmp_path / "source" / "main.py"
    source.parent.mkdir()
    source.write_text("import json\n", encoding="utf8")
    job = Routine("local", "Rotina local", str(source), schedule=Schedule(days=list(range(7))))
    config = Config(
        email="name@example.invalid",
        downloads=str(paths.root / "downloads"),
        logs=str(paths.root / "logs"),
        routines=[job],
    )
    save(paths, config)
    return paths, config, job


def test_configuration_cache_returns_independent_objects_and_detects_external_change(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    paths, _, _ = local_job(tmp_path)
    import orquestrador_app.config as module

    real = module.from_dict
    calls = []
    monkeypatch.setattr(module, "from_dict", lambda value: (calls.append(1), real(value))[1])
    one = load(paths)
    two = load(paths)
    one.email = "edited@example.invalid"
    assert two.email == "name@example.invalid" and len(calls) == 1
    save(paths, one)
    assert load(paths).email == one.email and len(calls) == 2
    with pytest.raises(ConfigurationConflict):
        save(paths, two)
    assert load(paths).email == one.email


def test_rare_job_latest_and_bounded_decision_state_ignore_global_recency(tmp_path: Path) -> None:
    paths, config, job = local_job(tmp_path)
    store = Store(paths.database)
    old = store.begin("old", job.id, "extraction", "2026-10-19T08:30")
    store.finish(old, "success", 0)
    store.succeed(job.id, old, period_for(job, NOW))
    with store.connect() as db:
        db.executemany(
            "INSERT INTO runs(id,batch,job,mode,started,status) VALUES(?, 'many','other','extraction',?,'success')",
            [(str(i), "2099-01-01T00:00:00+00:00") for i in range(1101)],
        )
        db.executemany(
            "INSERT INTO periods VALUES(?,?)", [(job.id, f"D:2000-01-{i}") for i in range(1, 29)]
        )
    assert store.latest(job.id)["id"] == old
    state = store.decision_state(job.id, period_for(job, NOW))
    assert len(state["periods"]) == len(state["occurrences"]) == 1
    assert Engine(paths, config).panel()[0]["last"]["id"] == old
    with pytest.raises(Exception):  # durable unique key still prevents duplicate attempts
        store.begin("another", job.id, "extraction", "2026-10-19T08:30")


def test_dependency_preview_is_conditional_and_does_not_write_runs_or_locks(tmp_path: Path) -> None:
    paths, config, job = local_job(tmp_path)
    child = Routine(
        "child",
        "Dependente",
        job.script,
        order=1,
        dependencies=[Dependency(job.id)],
        schedule=Schedule(days=list(range(7))),
    )
    config.routines.append(child)
    engine = Engine(paths, config)
    before = engine.store.history()
    preview = engine.preview([child.id], force=True)
    assert preview[0]["status"] == "blocked_dependency"
    preview = engine.preview([job.id, child.id], force=True)
    assert [r["status"] for r in preview] == ["simulation", "conditional"]
    assert preview[1]["requirements"] and engine.store.history() == before
    assert not paths.lock.exists()
    job.enabled = False
    assert engine.preview([job.id, child.id], force=True)[1]["blockers"]


def test_day_specific_times_pause_dates_and_single_recovery(tmp_path: Path) -> None:
    _, config, job = local_job(tmp_path)
    job.schedule.day_times = {"0": ["08:00"], "4": ["16:00"]}
    job.schedule.times = ["10:00"]
    assert eligible(job, {}, NOW).occurrence == "2026-10-19T08:00"
    assert not eligible(job, {"occurrences": {"2026-10-19T08:00"}}, NOW).due
    friday = NOW + dt.timedelta(days=4)
    assert next_execution(job, {}, friday).hour == 16
    job.schedule.paused_until = (NOW + dt.timedelta(hours=2)).isoformat()
    assert not eligible(job, {}, NOW).due
    assert eligible(job, {}, NOW + dt.timedelta(hours=2)).due
    job.schedule.excluded_dates = [NOW.date().isoformat()]
    assert not eligible(job, {}, NOW + dt.timedelta(hours=2)).due
    assert eligible(job, {}, NOW, force=True, manual=True).due
    config.calendar = {"name": "", "dates": ["2026-10-22"]}
    with pytest.raises(ValueError, match="calendário|Calendário"):
        validate(config)


def test_heartbeat_distinguishes_configured_active_stopped_stale_and_reused_pid(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    paths, config, _ = local_job(tmp_path)
    assert monitor_state(paths, config)["state"] == "configured"
    assert monitor_state(paths, config, installed=True)["state"] == "installed"
    write_heartbeat(paths, interval=30, state="idle")
    value = monitor_state(paths, config)
    assert value["active"] and value["communicating"]
    config.scheduling_paused = True
    assert monitor_state(paths, config)["state"] == "paused"
    assert (
        monitor_state(paths, config, now=dt.datetime.now(UTC) + dt.timedelta(minutes=3))["state"]
        == "no_communication"
    )
    write_heartbeat(paths, interval=30, state="stopped")
    assert not monitor_state(paths, config)["active"]
    import orquestrador_app.monitor_status as module

    monkeypatch.setattr(
        module.psutil,
        "Process",
        lambda _: SimpleNamespace(create_time=lambda: 0, is_running=lambda: True),
    )
    assert not monitor_state(paths, config)["active"]


def test_polling_is_slow_idle_and_fast_near_opportunity(tmp_path: Path) -> None:
    paths, config, job = local_job(tmp_path)
    store = Store(paths.database)
    job.schedule.times = ["10:00"]
    assert poll_seconds(config, store, NOW) == 30
    assert poll_seconds(config, store, NOW.replace(hour=9, minute=59, second=58)) == 2
    assert poll_seconds(config, store, NOW, running=True) == 1


def test_log_tail_preserves_multibyte_append_and_truncation(tmp_path: Path) -> None:
    file = tmp_path / "log"
    file.write_bytes(b"old\n\xc3")
    tail = LogTail()
    assert tail.read(file)["text"] == "old\n"
    with file.open("ab") as stream:
        stream.write(b"\xa7\nnew")
    update = tail.read(file)
    assert update["text"] == "ç\nnew" and not update["reset"]
    assert tail.read(file)["text"] == ""
    file.write_bytes(b"reset")
    assert tail.read(file) == {"text": "reset", "reset": True, "file": "log", "more": False}


def test_export_reconstructs_long_events_and_keeps_formula_literal(tmp_path: Path) -> None:
    text = "á😀" * 26000 + "\x01"
    target = tmp_path / "all.xlsx"
    export_data(
        target,
        {
            "Execuções": [{"id": "run", "value": "=1+1"}],
            "Eventos": [{"run": "run", "detail": text}],
        },
    )
    book = load_workbook(target, read_only=True, data_only=False)
    assert book.sheetnames[:2] == ["1-Execuções", "2-Eventos"]
    rows = list(book.worksheets[0].rows)
    assert rows[1][1].value == "=1+1" and rows[1][1].data_type == "s"
    overflow = list(book["Conteúdo longo"].values)[1:]
    assert "".join(row[5] for row in overflow) == text.replace("\x01", "\\u0001")
    book.close()


def test_events_conflicting_key_preserves_history(tmp_path: Path) -> None:
    target = tmp_path / "events.csv"
    publisher = Publisher(tmp_path / "backups", {"count": 5, "days": 90, "bytes": 100000})
    args = (["ID", "Value"], target, "run", ["ID"], ";", "utf-8-sig")
    publisher.merge_rows([{"ID": "1", "Value": "first"}], *args)
    before = target.read_bytes()
    assert publisher.merge_rows([{"ID": "1", "Value": "first"}], *args)["unchanged"]
    with pytest.raises(EventConflict) as raised:
        publisher.merge_rows([{"ID": "1", "Value": "different"}], *args)
    assert raised.value.conflicts and target.read_bytes() == before


def test_package_snapshot_dependencies_trash_and_safe_export(tmp_path: Path) -> None:
    paths, config, job = local_job(tmp_path)
    root = Path(job.script).parent
    job.package_root = str(root)
    helper = root / "helper.py"
    helper.write_text("import missing_caio_fixture_library\n")
    (root / "data.json").write_text('{"scale": 2}')
    (root / ".env").write_text("SECRET=notexported")
    check = check_dependencies(job)
    assert check["missing"] == ["missing_caio_fixture_library"] and not check["executed"]
    first = capture(paths, job)
    helper.write_text('print("changed")')
    capture(paths, job)
    restore(paths, config, job, first["sha256"])
    assert (Path(job.package_root) / "helper.py").read_text().startswith("import missing_")
    assert helper.read_text() == 'print("changed")'
    preview = export_routine_preview(config, job)
    archive = tmp_path / "routine.zip"
    export_routine(archive, preview)
    import zipfile

    with zipfile.ZipFile(archive) as bundle:
        assert not any(".env" in name for name in bundle.namelist())
    remove(paths, config, job.id)
    record = next((paths.root / "extras/trash").glob("*.json"))
    restore_registration(paths, load(paths), record)
    restored = load(paths)
    assert not restored.job(job.id).enabled
    assert restored.job(job.id).id == job.id


def test_notifications_retain_muted_messages_and_exact_run_links(tmp_path: Path) -> None:
    paths, config, job = local_job(tmp_path)
    store = Store(paths.database)
    assert enqueue(store, config, "failure", job.id, "one", "first")
    assert not enqueue(store, config, "failure", job.id, "two", "repeat")
    for _ in range(10):
        assert not enqueue(store, config, "failure", job.id, "two", "repeat")
    assert [r["run"] for r in claim(store)] == ["one"]
    assert claim(store) == []
    config.notifications["kinds"] = {"vpn": False}
    assert not enqueue(store, config, "vpn", job.id, "three", "vpn")
    assert {r["run"] for r in store.notification_history()} == {"one", "two", "three"}
    assert next(r for r in store.notification_history() if r["run"] == "two")["repeat_count"] == 11


def test_storage_preview_refuses_changed_files_and_keeps_unmanaged(tmp_path: Path) -> None:
    paths, config, _ = local_job(tmp_path)
    log = Path(config.logs) / "desktop.log"
    log.write_text("old")
    unknown = Path(config.logs) / "user.txt"
    unknown.write_text("keep")
    os.utime(log, (0, 0))
    preview = storage_preview(paths, config)
    assert len(preview["candidates"]) == 1
    log.write_text("changed")
    with pytest.raises(ValueError):
        clean_previewed(paths, preview)
    assert unknown.exists() and log.exists()


def test_quarantine_rejects_unowned_sources(tmp_path: Path) -> None:
    paths, _, _ = local_job(tmp_path)
    external = tmp_path / "external.xlsx"
    external.write_bytes(b"invalid")
    with pytest.raises(ValueError):
        quarantine(paths, external, "invalid", "run")
    owned = paths.root / "downloads/run-test/file.xlsx"
    owned.parent.mkdir()
    owned.write_bytes(b"invalid")
    record = quarantine(paths, owned, "invalid", "run")
    assert Path(record["file"]).read_bytes() == owned.read_bytes()
    assert record["sha256"] == digest(owned)


@pytest.mark.parametrize("mutate", [False, True])
def test_early_scheduler_probe_is_disabled_owned_and_always_cleaned(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, mutate: bool
) -> None:
    import orquestrador_app.tasks as module

    lifecycle, calls = [], []
    value: dict[str, str] = {}
    definition = SimpleNamespace(XmlText="")

    def register(name: str, definition: Any, *args: Any) -> None:
        assert args == (2, None, None, 3)
        value[name] = definition.XmlText
        if mutate:
            value[name] = value[name].replace("PT0S", "PT1S")
        calls.append("register")

    def delete(name: str, flags: int) -> None:
        del value[name]
        calls.append("delete")

    folder = SimpleNamespace(
        RegisterTaskDefinition=register,
        GetTask=lambda name: SimpleNamespace(Xml=value[name]),
        DeleteTask=delete,
    )
    service = SimpleNamespace(
        Connect=lambda: None, GetFolder=lambda _: folder, NewTask=lambda _: definition
    )
    client = SimpleNamespace(Dispatch=lambda _: service)
    package = ModuleType("win32com")
    package.client = client
    monkeypatch.setitem(sys.modules, "win32com", package)
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    monkeypatch.setitem(
        sys.modules,
        "pythoncom",
        SimpleNamespace(
            CoInitialize=lambda: lifecycle.append("in"),
            CoUninitialize=lambda: lifecycle.append("out"),
        ),
    )
    monkeypatch.setattr(module, "os", SimpleNamespace(name="nt", environ={"WINDIR": str(tmp_path)}))
    backend = WindowsTasks(Paths.user(tmp_path))
    monkeypatch.setattr(backend, "export", lambda name=None: value.get(name))
    if mutate:
        # Explicitly change a meaningful field while keeping ownership, to exercise cleanup on mismatch.
        original = register

        def changed(name: str, definition: Any, *args: Any) -> None:
            original(name, definition, *args)
            value[name] = value[name].replace("InteractiveToken", "S4U")

        folder.RegisterTaskDefinition = changed
        with pytest.raises(RuntimeError, match="divergiu"):
            backend.preflight("fixture")
    else:
        assert backend.preflight("fixture")["removed"]
    assert not value and calls == ["register", "delete"] and lifecycle == ["in", "out"]
    import xml.etree.ElementTree as ET

    root = ET.fromstring(definition.XmlText)
    assert root.findtext(f"{{{NS}}}Settings/{{{NS}}}Enabled") == "false"
    assert not list(root.find(f"{{{NS}}}Triggers"))


def test_history_model_emits_only_changed_row_and_pages_have_total(tmp_path: Path) -> None:
    app = QApplication.instance() or QApplication([])
    paths, _, job = local_job(tmp_path)
    store = Store(paths.database)
    for i in range(7):
        run = store.begin(str(i), job.id, "extraction", None)
        store.finish(run, "failed", 1, "driver")
    result = store.history_page(job=job.id, category="driver", size=3, page=1)
    assert result["total"] == 7 and len(result["rows"]) == 3
    model = HistoryModel()
    model.update(result["rows"], {job.id: job.label})
    changes = []
    resets = []
    model.dataChanged.connect(lambda *args: changes.append(1))
    model.modelReset.connect(lambda: resets.append(1))
    rows = [{**r} for r in result["rows"]]
    rows[1]["stage"] = "new"
    model.update(rows, {job.id: job.label})
    assert len(changes) == 1 and not resets
    app.processEvents()


def test_window_unsaved_guard_inline_validation_and_preferences(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    app = QApplication.instance() or QApplication([])
    paths = Paths.user(tmp_path / "window")
    paths.ensure()
    save(paths, defaults(paths, source_root()))
    window = MainWindow(paths)
    settle(window)
    window.timer.stop()
    window.dev_panel.timer.stop()
    window.show()
    try:
        window.navigation.setCurrentRow(2)
        window.common_settings.email.setText("invalid")
        assert window.settings_dirty() and not window.save_settings()
        assert window.common_settings.email.text() == "invalid" and window.settings_error.text()
        monkeypatch.setattr(
            QMessageBox, "question", lambda *args: QMessageBox.StandardButton.Cancel
        )
        window.navigation.setCurrentRow(0)
        assert window.tabs.currentIndex() == 2
        window.reset_settings(load(paths))
        window.follow_log.setChecked(False)
        window.save_window_preferences()
        assert read_json(paths.root / "state/window-preferences.json")["follow_log"] is False
        window.panel.selectRow(1)
        selected = window.selected_ids()
        window.refresh()
        settle(window)
        assert window.selected_ids() == selected
    finally:
        window.exit_requested = True
        window.close()
        settle(window)
        app.processEvents()


def test_locked_database_observation_keeps_qt_controls_responsive(tmp_path: Path) -> None:
    import sqlite3
    import time

    from PySide6.QtCore import QTimer

    app = QApplication.instance() or QApplication([])
    paths = Paths.user(tmp_path / "locked")
    paths.ensure()
    save(paths, defaults(paths, source_root()))
    window = MainWindow(paths)
    settle(window)
    window.timer.stop()
    window.dev_panel.timer.stop()
    with window.store.connect() as db:
        db.execute(
            "INSERT INTO notifications(id,at,kind,job,run,message) VALUES('pending-lock-test',?,'failure',?,'','fixture')",
            (NOW.isoformat(), window.config.routines[0].id),
        )
    connection = sqlite3.connect(paths.database)
    connection.execute("BEGIN IMMEDIATE")
    ticks = []
    timer = QTimer()
    timer.timeout.connect(lambda: ticks.append(1))
    timer.start(10)
    try:
        window.refresh()
        deadline = time.monotonic() + 0.12
        while time.monotonic() < deadline:
            app.processEvents()
            time.sleep(0.001)
        assert len(ticks) >= 5 and window.read_task is not None
    finally:
        connection.rollback()
        connection.close()
        timer.stop()
        settle(window)
        window.exit_requested = True
        window.close()
        app.processEvents()


def test_routine_dialog_keeps_revision_from_start_of_edit(tmp_path: Path) -> None:
    QApplication.instance() or QApplication([])
    paths, _, _ = local_job(tmp_path)
    initial = load(paths)
    dialog = RoutineDialog(initial, initial.routines[0])
    newer = load(paths)
    newer.email = "changed@example.invalid"
    save(paths, newer)
    try:
        edited = dialog.collect_values()
        with pytest.raises(ConfigurationConflict):
            save(paths, edited)
        assert load(paths).email == newer.email
    finally:
        dialog.done(0)
