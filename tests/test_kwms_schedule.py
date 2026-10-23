from __future__ import annotations

import datetime as dt
import json
import logging
from pathlib import Path
import sys
import tempfile
import unittest
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import Workbook, load_workbook
import automation_schedule as scheduling
import kwms_alocacoes_tool_edge as kwms
import pipefy_orchestrator as orchestration
import update_installation as updater


class FrequencyTests(unittest.TestCase):
    def test_daily_runs_once_per_local_date_after_its_time(self):
        schedule = scheduling.parse_schedule({"frequency": "daily", "time": "08:30"})
        before = dt.datetime.fromisoformat("2026-10-21T08:29:59-03:00")
        now = dt.datetime.fromisoformat("2026-10-21T10:30:00-03:00")
        self.assertFalse(scheduling.due_reason(schedule, None, before)[0])
        self.assertTrue(scheduling.due_reason(schedule, "2026-10-20T12:00:00-03:00", now)[0])
        self.assertFalse(scheduling.due_reason(schedule, "2026-10-21T08:31:00-03:00", now)[0])
        self.assertTrue(scheduling.due_reason(schedule, "2026-10-21T01:00:00+00:00", now)[0])

    def test_weekly_catches_up_and_does_not_repeat_when_year_changes(self):
        schedule = scheduling.parse_schedule({"frequency": "weekly", "weekday": "Monday", "weekdays_only": True})
        friday = dt.datetime.fromisoformat("2027-01-01T10:30:00-03:00")
        monday = dt.datetime.fromisoformat("2027-01-04T08:30:00-03:00")
        self.assertTrue(scheduling.due_reason(schedule, None, friday)[0])
        self.assertFalse(scheduling.due_reason(schedule, "2026-12-28T08:40:00-03:00", friday)[0])
        self.assertTrue(scheduling.due_reason(schedule, "2027-01-01T08:40:00-03:00", monday)[0])
        sunday = dt.datetime.fromisoformat("2027-01-03T12:00:00-03:00")
        self.assertFalse(scheduling.due_reason(schedule, None, sunday)[0])

    def test_weekly_waits_for_chosen_weekday(self):
        schedule = scheduling.parse_schedule({"frequency": "weekly", "weekday": "Friday", "time": "10:30"})
        wednesday = dt.datetime.fromisoformat("2026-10-21T16:30:00-03:00")
        friday = dt.datetime.fromisoformat("2026-10-23T10:30:00-03:00")
        self.assertFalse(scheduling.due_reason(schedule, None, wednesday)[0])
        self.assertTrue(scheduling.due_reason(schedule, None, friday)[0])

    def test_invalid_schedule_and_corrupt_state_are_not_silently_ignored(self):
        for raw in ({"frequency": "daliy"}, {"time": "8:30"}, {"weekday": "Monday", "unknown": True},
                    {"frequency": "weekly", "weekday": "Saturday", "weekdays_only": True}):
            with self.subTest(raw=raw), self.assertRaises(scheduling.ScheduleError):
                scheduling.parse_schedule(raw)
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "state.json"
            path.write_text('{"broken"', encoding="utf-8")
            with self.assertRaises(scheduling.ScheduleError):
                scheduling.read_state(path)

    def test_preferences_preserve_unrelated_config_and_make_backup(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            config_path = root / "config.json"
            original = {"custom": {"keep": True}, "schedules": ["08:30", "16:30"],
                        "job_schedules": {"other": {"frequency": "weekly"}}}
            config_path.write_text(json.dumps(original), encoding="utf-8")
            job = mock.Mock(name="job")
            job.name = "kwms"
            scheduling.save_preferences(config_path, [job], {"kwms": {"enabled": True, "frequency": "weekly",
                                      "weekday": "Tuesday", "time": "10:30", "weekdays_only": True}})
            result = json.loads(config_path.read_text(encoding="utf-8"))
            self.assertEqual(result["custom"], original["custom"])
            self.assertEqual(result["job_schedules"]["other"], original["job_schedules"]["other"])
            self.assertTrue(result["schedule_setup_completed"])
            self.assertEqual(json.loads(config_path.with_suffix(".json.bak").read_text()), original)


class LocalOrchestrationTests(unittest.TestCase):
    def make_orchestrator(self, root, jobs):
        config = {"schema_version": 1, "project_root": ".", "runtime_dir": "state", "python_executable": sys.executable,
                  "inter_job_delay_seconds": 300, "driver": {"managed_path": "driver/msedgedriver.exe"}, "jobs": jobs}
        path = root / "config.json"
        path.write_text(json.dumps(config), encoding="utf-8")
        return orchestration.Orchestrator(orchestration.load_settings(path), logging.getLogger("kwms_test"))

    def test_vpn_skip_continues_even_when_continue_on_error_is_false_without_extra_retries(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "vpn.py").write_text("from pathlib import Path\np=Path('vpn-count.txt')\np.write_text(p.read_text()+'x' if p.exists() else 'x')\nraise SystemExit(20)\n")
            (root / "next.py").write_text("from pathlib import Path\nPath('next-ran.txt').write_text('ok')\n")
            jobs = [{"name": "kwms", "script": "vpn.py", "command": "", "retries": 3,
                     "skip_exit_codes": [20], "schedule": {"frequency": "weekly", "time": "00:00"}},
                    {"name": "next", "script": "next.py", "command": "", "retries": 0}]
            runner = self.make_orchestrator(root, jobs)
            with mock.patch.object(orchestration.time, "sleep") as sleep:
                code = runner.run(None, None, True, False)
                self.assertNotIn(mock.call(300), sleep.call_args_list)
                self.assertNotIn(mock.call(30), sleep.call_args_list)
            self.assertEqual(code, 0)
            self.assertEqual((root / "vpn-count.txt").read_text(), "x")
            self.assertTrue((root / "next-ran.txt").exists())
            summary = json.loads((root / "state" / "last_run.json").read_text())
            self.assertTrue(summary["has_warnings"])
            self.assertEqual([entry["status"] for entry in summary["jobs"]], ["skipped", "success"])
            state = scheduling.read_state(root / "state" / "job_schedule_state.json")
            self.assertNotIn("kwms", state["jobs"])

    def test_daily_success_is_persisted_and_skips_browser_check_on_restart(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "job.py").write_text("from pathlib import Path\np=Path('count.txt')\np.write_text(p.read_text()+'x' if p.exists() else 'x')\n")
            jobs = [{"name": "daily", "script": "job.py", "command": "", "retries": 0,
                     "schedule": {"frequency": "daily", "time": "00:00"}}]
            runner = self.make_orchestrator(root, jobs)
            self.assertEqual(runner.run(None, 0, True, False), 0)
            runner = self.make_orchestrator(root, jobs)
            with mock.patch.object(runner.driver_guard, "ensure") as guard:
                self.assertEqual(runner.run(None, 0, False, False), 0)
                guard.assert_not_called()
            self.assertEqual((root / "count.txt").read_text(), "x")
            self.assertEqual(runner.run(None, 0, True, False, force=True), 0)
            self.assertEqual((root / "count.txt").read_text(), "xx")

    def test_update_preserves_existing_json_extra_extractor_and_previous_code(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, target = root / "update", root / "installation"
            source.mkdir()
            (target / "scripts").mkdir(parents=True)
            existing = {"schema_version": 1, "project_root": "scripts", "runtime_dir": "../state",
                        "driver": {"managed_path": "my-driver.exe"}, "schedules": ["09:00"], "custom": "keep",
                        "job_schedules": {"custom": {"frequency": "weekly"}}, "jobs": []}
            incoming = {"schema_version": 1, "schedule_defaults": {"frequency": "daily"},
                        "job_schedules": {"kwms": {"frequency": "weekly"}},
                        "jobs": [{"name": "kwms", "script": "kwms.py"}]}
            (source / updater.CONFIG_NAME).write_text(json.dumps(incoming))
            (target / updater.CONFIG_NAME).write_text(json.dumps(existing))
            (source / "app.py").write_text("new")
            (target / "app.py").write_text("old")
            (target / "scripts" / "pipefy_status_atms_tool_edge.py").write_text("extra")
            backup = updater.update_installation(source, target)
            result = json.loads((target / updater.CONFIG_NAME).read_text())
            self.assertEqual(result["driver"], existing["driver"])
            self.assertEqual(result["schedules"], ["09:00"])
            self.assertEqual(result["custom"], "keep")
            self.assertEqual(result["job_schedules"]["custom"], {"frequency": "weekly"})
            self.assertEqual(result["jobs"][0]["name"], "kwms")
            self.assertEqual((backup / "app.py").read_text(), "old")
            self.assertEqual((target / "app.py").read_text(), "new")
            self.assertEqual((target / "scripts" / "pipefy_status_atms_tool_edge.py").read_text(), "extra")


class KwmsTests(unittest.TestCase):
    def workbook(self, path, rows=None):
        book = Workbook()
        for row in rows or [["Engagement", "Horas"], [12345, 8]]:
            book.active.append(row)
        book.save(path)
        book.close()

    def test_vpn_error_recognized_from_browser_exception_and_rendered_page(self):
        for detail in ("unknown error: net::ERR_EMPTY_RESPONSE", "kwms.br.kworld.kpmg.com didn’t send any data."):
            self.assertTrue(kwms.is_empty_response(detail))
        self.assertFalse(kwms.is_empty_response("Timeout de login"))
        driver = mock.Mock()
        driver.execute_script.return_value = "ERR_EMPTY_RESPONSE"
        with self.assertRaises(kwms.VpnRequiredError):
            kwms.check_vpn_error(driver)
        driver.execute_script.return_value = ""
        with self.assertRaises(kwms.VpnRequiredError):
            kwms.check_vpn_error(driver, RuntimeError("net::ERR_EMPTY_RESPONSE"))

    def test_exactly_two_vpn_attempts_separated_by_30_seconds(self):
        wait = mock.Mock()
        operation = mock.Mock(side_effect=[kwms.VpnRequiredError(), "done"])
        self.assertEqual(kwms.with_vpn_retry(operation, wait), "done")
        self.assertEqual(operation.call_count, 2)
        wait.assert_called_once_with(30)
        wait.reset_mock()
        operation = mock.Mock(side_effect=kwms.VpnRequiredError())
        with self.assertRaises(kwms.VpnRequiredError):
            kwms.with_vpn_retry(operation, wait)
        self.assertEqual(operation.call_count, 2)
        wait.assert_called_once_with(30)

    def test_other_errors_do_not_request_vpn_or_retry(self):
        operation = mock.Mock(side_effect=ValueError("Invalid workbook"))
        wait = mock.Mock()
        with self.assertRaises(ValueError):
            kwms.with_vpn_retry(operation, wait)
        self.assertEqual(operation.call_count, 1)
        wait.assert_not_called()

    def test_publication_uses_exact_filename_and_replaces_old_content(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            self.workbook(root / "download.xlsx")
            destination = root / "KWMS"
            destination.mkdir()
            old = destination / kwms.FINAL_FILENAME
            old.write_bytes(b"old")
            target = kwms.publish_workbook(root / "download.xlsx", destination)
            self.assertEqual(target.name, "Detalhes de Alocação_- Planos Aprovados e Baixados.xlsx")
            book = load_workbook(target, read_only=True)
            self.assertEqual(book.active["A2"].value, 12345)
            book.close()
            self.assertEqual(list(destination.iterdir()), [target])

    def test_corrupt_empty_or_locked_download_does_not_replace_previous_file(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            destination = root / "KWMS"
            destination.mkdir()
            old = destination / kwms.FINAL_FILENAME
            old.write_bytes(b"previous workbook")
            source = root / "download.xlsx"
            source.write_bytes(b"<html>error</html>")
            with self.assertRaises(ValueError):
                kwms.publish_workbook(source, destination)
            self.assertEqual(old.read_bytes(), b"previous workbook")
            self.workbook(source, [["Header"]])
            with self.assertRaises(ValueError):
                kwms.publish_workbook(source, destination)
            self.workbook(source)
            with mock.patch.object(kwms.os, "replace", side_effect=PermissionError("Excel lock")):
                with self.assertRaises(PermissionError):
                    kwms.publish_workbook(source, destination)
            self.assertEqual(old.read_bytes(), b"previous workbook")
            self.assertEqual(list(destination.iterdir()), [old])

    def test_download_accepts_timestamp_and_waits_for_partial_file(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            name = "Detalhes de Alocação_- Planos Aprovados e Baixados - 1792430784969.xlsx"
            target = root / name
            target.write_bytes(b"workbook")
            partial = root / (name + ".crdownload")
            partial.write_bytes(b"unfinished")
            driver = mock.Mock()
            driver.execute_script.return_value = "Relatório"
            with mock.patch.object(kwms.time, "sleep", side_effect=lambda seconds: partial.unlink(missing_ok=True)):
                self.assertEqual(kwms.wait_for_download(root, 3, driver), target)
            self.assertFalse(kwms.matches_download_name("Outro relatório.xlsx"))

    def test_kwms_exit_20_and_browser_is_closed_after_second_vpn_failure(self):
        with tempfile.TemporaryDirectory() as temporary:
            args = kwms.build_parser().parse_args(["--work-dir", temporary, "run"])
            driver = mock.Mock()
            with mock.patch.object(kwms, "build_edge", return_value=driver), \
                 mock.patch.object(kwms, "open_reporting", side_effect=kwms.VpnRequiredError()) as open_page, \
                 mock.patch.object(kwms, "notify_vpn_and_wait") as wait:
                self.assertEqual(kwms.execute(args), 20)
            self.assertEqual(open_page.call_count, 2)
            wait.assert_called_once_with(30)
            driver.quit.assert_called_once()


if __name__ == "__main__":
    unittest.main()
